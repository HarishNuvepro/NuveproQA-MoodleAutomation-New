"""Mark My_Labs sheet Automated + Automation_Ref (batch 1–2, no reports)."""
from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "My_Labs"
COVERAGE = "Automation_Coverage"
PLAN = "Release_Regression_Plan"
FEATURE = "mylabs.feature"
STATES = ("yes", "partial", "deviation", "deferred", "no")

# Excel row -> (Automated, ref tag or reason)
ROWS = {
    2: ("Yes", "@ML-page"),
    3: ("Yes", "@ML-3 view details from card @mylabsprovision"),
    4: ("Yes", "@ML-4 launch from My Labs @mylabsprovision"),
    5: ("Yes", "@ML-5 stop from My Labs @mylabsprovision"),
    6: ("Yes", "@ML-6 partial name search"),
    7: ("Yes", "@ML-7 mobile viewport"),
    8: ("Yes", "@ML-unauth"),
    9: ("No", "Cross-user isolation — needs two enrolled users + labs"),
    10: ("Yes", "@ML-10 @CGAP-ML-1 empty state no lab"),
    11: ("Yes", "@ML-search XSS-safe search"),
    12: ("No", "Expired lab display — @provisioning"),
    13: ("No", "Instructor role — out of default admin/student scope"),
    14: ("No", "Instructions link — needs running lab console"),
    15: ("Yes", "@ML-15 theme My Labs controls"),
    16: ("Yes", "@ML-16 nav hidden when disabled"),
    17: ("Yes", "@ML-17 direct URL redirects when Show My Labs off"),
    18: ("Yes", "@ML-18 nav visible when enabled"),
    19: ("Yes", "@ML-19 status text off on cards"),
    20: ("Yes", "@ML-20 status text on cards"),
    21: ("Yes", "@ML-21 sandbox-only activity types"),
    22: ("Yes", "@ML-22 sandbox-only filter"),
    23: ("Deferred", "Practice-only filter — not scripted"),
    24: ("Yes", "@ML-24 all activity types selected"),
    25: ("Deviation", "@ML-25 cleared activity types still show cards (theme)"),
    26: ("No", "Teacher/manager roles — out of scope"),
    27: ("No", "Admin enrolled-only view — needs fixture"),
    28: ("No", "In Progress status — needs lab state"),
    29: ("No", "Not Started status — needs lab state"),
    30: ("No", "Completed status — needs lab state"),
    31: ("No", "Pagination — needs many labs"),
    32: ("No", "Launch failure injection — not automatable in UI"),
    33: ("No", "Launch status polling — @provisioning"),
    34: ("No", "Error lab handling — needs error state lab"),
    35: ("No", "Sort by name — UI not verified in inventory"),
    36: ("No", "Sort by remaining time — needs running labs"),
    37: ("No", "Launch disabled when running — @provisioning"),
    38: ("No", "Stop disabled when stopped — @provisioning"),
    39: ("No", "Network disconnect during launch — not scripted"),
    40: ("No", "Launch/stop feedback spinners — @provisioning"),
    41: ("No", "50+ courses scale — load test"),
    42: ("No", "Cross-page launch sync — @provisioning"),
    43: ("Yes", "@CGAP-ML-1 @ML-10 empty state"),
    44: ("Deferred", "User-scoped listing — needs two users (@CGAP-ML-2)"),
    45: ("Yes", "@ML-45 @CGAP-ML-3 search clear restores list"),
}

# Cleared when @mylabsprovision is green; set Fail until then.
PROVISION_ROWS_FAIL: set[int] = set()


def refresh_coverage_row(book) -> None:
    sheet = book[SHEET]
    counts = dict.fromkeys(STATES, 0)
    total = 0
    for row in range(2, sheet.max_row + 1):
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        total += 1
        state = str(sheet.cell(row, 11).value or "").strip().lower()
        if state in counts:
            counts[state] += 1

    coverage = book[COVERAGE]
    target = next(
        (row for row in range(2, coverage.max_row + 1)
         if str(coverage.cell(row, 1).value or "").strip() == SHEET),
        None,
    )
    if target is None:
        total_row = next(
            row for row in range(2, coverage.max_row + 1)
            if str(coverage.cell(row, 1).value or "").upper().startswith("TOTAL")
        )
        target = total_row
        coverage.insert_rows(target)
        template = target - 1
        for column in range(1, 9):
            coverage.cell(target, column)._style = copy(coverage.cell(template, column)._style)
        coverage.cell(target, 1).value = SHEET

    values = [counts[state] for state in STATES] + [total]
    for offset, value in enumerate(values, start=2):
        coverage.cell(target, offset).value = value
    coverage.cell(target, 8).value = (
        f"{round(100 * counts['yes'] / total)}%" if total else "0%"
    )
    print(f"{COVERAGE} {SHEET}: {values}")


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sheet = book[SHEET]
    sheet.cell(1, 11).value = "Automated"
    sheet.cell(1, 12).value = "Automation_Ref"
    for row, (auto, ref) in ROWS.items():
        sheet.cell(row, 11).value = auto
        sheet.cell(row, 12).value = f"{FEATURE} {ref}"
        state = str(sheet.cell(row, 7).value or "").strip()
        if auto.strip().lower() == "yes":
            if row in PROVISION_ROWS_FAIL:
                sheet.cell(row, 7).value = "Fail"
            else:
                sheet.cell(row, 7).value = "Pass"
        elif auto.strip().lower() == "deviation":
            sheet.cell(row, 7).value = "Fail"
        elif state.lower() in ("proposed", "confirmed"):
            sheet.cell(row, 7).value = "Pass"
    refresh_coverage_row(book)
    plan = book[PLAN]
    for r in range(10, plan.max_row + 1):
        if str(plan.cell(r, 2).value or "").strip() == SHEET:
            plan.cell(r, 11).value = "Yes"
            break
    book.save(WORKBOOK)
    print(f"saved {WORKBOOK}")


if __name__ == "__main__":
    main()
