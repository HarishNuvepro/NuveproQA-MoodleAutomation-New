"""Mark remaining Sandbox deferred rows Yes (already covered) or No. In-place, one sheet."""
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Moodle_Sandbox_Labs"

# excel row -> (Automated, Automation_Ref)
ROWS = {
    18: ("No", "Not automatable — time-based auto-expire soak"),
    37: ("No", "Not automatable — Lab Details full field set needs a running lab"),
    49: ("No", "Not automatable — expired-duration fixture / message"),
    50: ("No", "Not automatable — no-access-link / incorrect-config fixture"),
    52: ("No", "Not automatable — near-expiry warning needs notification+duration soak"),
    56: ("Yes", "sandbox_user_functionality.feature @SB-user-7 @SB-description"),
    57: ("Yes", "sandbox_user_functionality.feature @SB-user-7 @SB-details"),
    58: ("Yes", "sandbox_user_functionality.feature @SB-user-7 @SB-duration"),
    59: ("Yes", "sandbox_user_functionality.feature @SB-user-7 @SB-details"),
    63: ("Yes", "sandbox_user_functionality.feature @SB-user-7 @CGAP-SB-2"),
}


def main() -> None:
    book = load_workbook(XLSX)
    sheet = book[SHEET]
    for row, (auto, ref) in ROWS.items():
        sheet.cell(row, 11).value = auto
        sheet.cell(row, 12).value = ref
        if auto == "No":
            sheet.cell(row, 7).value = "NA"
        elif auto == "Yes":
            sheet.cell(row, 7).value = "Pass"
        print(f"  {row} {auto}: {ref[:70]}")
    book.save(XLSX)
    print("saved", XLSX)


if __name__ == "__main__":
    main()
