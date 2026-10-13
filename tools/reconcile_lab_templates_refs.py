"""Phase 0: align Moodle_Lab_Templates Yes rows with scenario tags that exist in features."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Moodle_Lab_Templates"
FEATURES = Path("src/test/resources/features")

TAG_RE = re.compile(r"@LT-(\d+)|@CGAP-LT-(\d+)")


def collect_tags() -> set[str]:
    found: set[str] = set()
    for path in FEATURES.glob("*.feature"):
        text = path.read_text(encoding="utf-8")
        for m in TAG_RE.finditer(text):
            if m.group(1):
                found.add(f"LT-{m.group(1)}")
            elif m.group(2):
                found.add(f"CGAP-LT-{m.group(2)}")
        for m in re.finditer(r"@LT-role\b", text):
            found.add("LT-role")
    return found


def primary_tag(ref: str) -> str | None:
    m = TAG_RE.search(ref or "")
    if not m:
        return None
    if m.group(1):
        return f"LT-{m.group(1)}"
    if m.group(2):
        return f"CGAP-LT-{m.group(2)}"
    return None


def main() -> None:
    live = collect_tags()
    print(f"scenario tags in repo: {len(live)}")
    book = openpyxl.load_workbook(WORKBOOK)
    sheet = book[SHEET]
    downgraded = 0
    for row in range(2, sheet.max_row + 1):
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        auto = str(sheet.cell(row, 11).value or "").strip().lower()
        if auto != "yes":
            continue
        ref = str(sheet.cell(row, 12).value or "")
        tag = primary_tag(ref)
        if tag is None or tag not in live:
            sheet.cell(row, 11).value = "Partial"
            note = f"Phase 0 — no Cucumber @{tag or '?'} yet; was: {ref[:80]}"
            sheet.cell(row, 12).value = note
            if str(sheet.cell(row, 7).value or "").strip().lower() == "pass":
                sheet.cell(row, 7).value = "NA"
            downgraded += 1
            print(f"row {row} -> Partial ({tag})")
    from update_lab_templates_coverage import refresh_coverage_row

    refresh_coverage_row(book)
    book.save(WORKBOOK)
    print(f"downgraded {downgraded} false-Yes Lab Template rows; Automation_Coverage refreshed")


if __name__ == "__main__":
    main()
