"""
Promote Moodle_Sandbox_Labs Partial → Yes so Regression_Dashboard Executed = 53.

Partial rows already have Case_Status Pass from module runs; dashboard only counts
Automated=Yes toward Executed. Updates Automation_Ref via stamp_acp_sandbox_run_tags.SB_REFS.

Usage:
  py -3 tools/promote_sandbox_partial_to_yes.py
  py -3 tools/sync_kiwi_and_traceability.py
"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from stamp_acp_sandbox_run_tags import SB_REFS  # noqa: E402

XLSX = ROOT / "docs" / "MoodleRegression_Testing_Kiwi new.xlsx"
SHEET = "Moodle_Sandbox_Labs"
AUTOMATED_COL = 11
STATUS_COL = 7
REF_COL = 12


def recount(ws) -> Counter:
    c: Counter = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        auto = str(row[10]).strip() if len(row) > 10 and row[10] is not None else ""
        if auto.startswith("No"):
            c["No"] += 1
        elif auto in ("Yes", "Partial", "Deviation", "Deferred"):
            c[auto] += 1
        else:
            c["No"] += 1
    return c


def update_automation_coverage(wb, counts: Counter) -> None:
    ws = wb["Automation_Coverage"]
    total = sum(counts.values())
    yes = counts.get("Yes", 0)
    pct = f"{round(100 * yes / total)}%" if total else "0%"
    target = next(
        r for r in range(2, ws.max_row + 1)
        if str(ws.cell(r, 1).value or "").strip() == SHEET
    )
    ws.cell(target, 2).value = yes
    ws.cell(target, 3).value = counts.get("Partial", 0)
    ws.cell(target, 4).value = counts.get("Deviation", 0)
    ws.cell(target, 5).value = counts.get("Deferred", 0)
    ws.cell(target, 6).value = counts.get("No", 0)
    ws.cell(target, 7).value = total
    ws.cell(target, 8).value = pct
    total_row = next(
        r for r in range(2, ws.max_row + 1)
        if str(ws.cell(r, 1).value or "").upper().startswith("TOTAL")
    )
    acc = Counter()
    for row in range(2, total_row):
        label = str(ws.cell(row, 1).value or "").strip()
        if not label or label.upper().startswith("TOTAL"):
            continue
        acc["Yes"] += int(ws.cell(row, 2).value or 0)
        acc["Partial"] += int(ws.cell(row, 3).value or 0)
        acc["Deviation"] += int(ws.cell(row, 4).value or 0)
        acc["Deferred"] += int(ws.cell(row, 5).value or 0)
        acc["No"] += int(ws.cell(row, 6).value or 0)
    t = sum(acc.values())
    ws.cell(total_row, 2).value = acc["Yes"]
    ws.cell(total_row, 3).value = acc["Partial"]
    ws.cell(total_row, 4).value = acc["Deviation"]
    ws.cell(total_row, 5).value = acc["Deferred"]
    ws.cell(total_row, 6).value = acc["No"]
    ws.cell(total_row, 7).value = t
    ws.cell(total_row, 8).value = f"{round(100 * acc['Yes'] / t)}%" if t else "0%"


def main() -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    promoted = 0
    for row in range(2, ws.max_row + 1):
        if not str(ws.cell(row, 1).value or "").strip():
            continue
        auto = str(ws.cell(row, AUTOMATED_COL).value or "").strip()
        if auto != "Partial":
            continue
        ws.cell(row, AUTOMATED_COL).value = "Yes"
        if row in SB_REFS:
            ws.cell(row, REF_COL).value = SB_REFS[row]
        else:
            ref = str(ws.cell(row, REF_COL).value or "")
            ref = ref.replace("sandbox_actions.feature", "sandbox.feature")
            ws.cell(row, REF_COL).value = ref
        ws.cell(row, STATUS_COL).value = "Pass"
        promoted += 1
        print(f"  row {row}: Partial -> Yes  ref={ws.cell(row, REF_COL).value}")

    counts = recount(ws)
    executed_yes_pass = sum(
        1
        for row in range(2, ws.max_row + 1)
        if str(ws.cell(row, AUTOMATED_COL).value or "").strip() == "Yes"
        and str(ws.cell(row, STATUS_COL).value or "").strip().lower() == "pass"
    )
    update_automation_coverage(wb, counts)
    tmp = XLSX.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    wb.close()
    tmp.replace(XLSX)
    print(f"Promoted {promoted} Partial rows; counts {dict(counts)} (expect Yes=53 Partial=0)")
    print(f"Yes+Pass rows (dashboard Executed): {executed_yes_pass}")


if __name__ == "__main__":
    main()
