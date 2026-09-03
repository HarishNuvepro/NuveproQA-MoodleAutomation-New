"""
Put All Course Page on Regression_Dashboard directly below Moodle_Sandbox_Labs.

- Case_Status = Pass for every Automated=Yes row (same convention as Moodle_Sandbox_Labs),
  so the dashboard Executed formula counts them.
- Rewrite the module block so All Course Page sits in the QA-inserted slot and appears ONCE.
  The old row referenced 'All Course Page' without its trailing space, so its formulas were #REF.
- Updates the workbook in place. No backup copies.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "MoodleRegression_Testing_Kiwi new.xlsx"
SHEET = "All Course Page "
DASH = "Regression_Dashboard"
ANCHOR = "Moodle_Sandbox_Labs"
FIRST_ROW = 5


def sheet_ref(name: str) -> str:
    return f"'{name}'" if any(ch in name for ch in " &-") else name


def executed_formula(ref: str) -> str:
    return "".join(
        f'{"=" if i == 0 else "+"}COUNTIFS({ref}!K2:K2000,"Yes",{ref}!G2:G2000,"{status}")'
        for i, status in enumerate(("Pass", "Fail", "Failed", "Blocked"))
    )


def mark_pass(wb) -> int:
    ws = wb[SHEET]
    changed = 0
    for r in range(2, ws.max_row + 1):
        if not ws.cell(r, 1).value:
            continue
        if str(ws.cell(r, 11).value).strip() == "Yes":
            if str(ws.cell(r, 7).value).strip().lower() != "pass":
                changed += 1
            ws.cell(r, 7).value = "Pass"
    return changed


def rewrite_dashboard(wb) -> None:
    ws = wb[DASH]
    rows: list[list] = []
    total_row = None
    for r in range(FIRST_ROW, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name:
            continue
        if str(name).strip().upper().startswith("TOTAL"):
            total_row = [ws.cell(r, c).value for c in range(1, 9)]
            continue
        if str(name).strip() == SHEET.strip():
            continue  # stale row; re-added at the anchor below
        rows.append([ws.cell(r, c).value for c in range(1, 9)])

    ref = sheet_ref(SHEET)
    acp = [17, SHEET.strip(),
           f"=COUNTA({ref}!A2:A2000)", None,
           executed_formula(ref),
           f'=COUNTIF({ref}!K2:K2000,"No")',
           None, None]
    at = next((i for i, row in enumerate(rows)
               if str(row[1]).strip() == ANCHOR), len(rows) - 1)
    rows.insert(at + 1, acp)

    last_used = ws.max_row
    for i, row in enumerate(rows):
        r = FIRST_ROW + i
        for c in range(1, 7):
            ws.cell(r, c).value = row[c - 1]
        # Not Executed / Completion % are row-local, so regenerate them after the move.
        if isinstance(row[6], str) or row[6] is None:
            ws.cell(r, 7).value = f"=max(0, C{r}-E{r}-F{r})"
            ws.cell(r, 8).value = f"=if((C{r}-F{r})<=0, 0, E{r}/(C{r}-F{r}))"
        else:
            ws.cell(r, 7).value = row[6]
            ws.cell(r, 8).value = row[7]

    total_at = FIRST_ROW + len(rows)
    ws.cell(total_at, 1).value = total_row[0] if total_row else None
    ws.cell(total_at, 2).value = "TOTAL"
    ws.cell(total_at, 3).value = f"=sum(C{FIRST_ROW}:C{total_at - 1})"
    for c, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
        ws.cell(total_at, c).value = (
            f"=sum({col}{FIRST_ROW}:{col}{total_at - 1})" if c != 4 else None)
    ws.cell(total_at, 8).value = (
        f"=if((C{total_at}-F{total_at})<=0, 0, E{total_at}/(C{total_at}-F{total_at}))")

    for r in range(total_at + 1, last_used + 1):
        for c in range(1, 9):
            ws.cell(r, c).value = None
    print(f"{DASH}: All Course Page at row {FIRST_ROW + at + 1}, TOTAL at row {total_at}")


def main() -> None:
    wb = load_workbook(XLSX)
    changed = mark_pass(wb)
    rewrite_dashboard(wb)
    tmp = XLSX.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    wb.close()
    tmp.replace(XLSX)
    print("Saved", XLSX)
    print("Case_Status set to Pass on", changed, "newly executed rows")


if __name__ == "__main__":
    main()
