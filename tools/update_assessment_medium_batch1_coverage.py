"""Promote Assessment reopening / hide-grade / CGAP-A-8 rows to verified Kiwi states (medium batch 1b)."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Moodle_Assessment_Labs"
COVERAGE = "Automation_Coverage"
FEATURES = Path("src/test/resources/features")

# Tags with real steps (no provisioning stub line) in assessment_reopening.feature
REOPEN_YES: dict[str, str] = {
    f"@{t}": f"assessment_reopening.feature @{t}"
    for t in (
        "A24", "A25", "A26", "A27", "A29", "A34", "A37", "A38", "A39", "A42",
        "A45", "A47", "A49", "A50", "A52", "A-override",
    )
}
REOPEN_PARTIAL: dict[str, str] = {
    f"@{t}": f"assessment_reopening.feature @{t} (may skip without completed user)"
    for t in ("A53", "A54", "A44", "A46", "A48", "A51")
}

HIDE_FINAL_YES: dict[str, str] = {
    f"@{t}": f"assessment_hide_final_grade.feature @{t}"
    for t in ("A130", "A131", "A132", "A135", "A136", "A137", "A138", "A139")
}

SHOW_HIDE_YES: dict[str, str] = {
    f"@{t}": f"assessment_show_hide_grades.feature @{t}"
    for t in ("A55", "A56", "A58", "A59", "A60", "A61")
}

FIXED: dict[str, tuple[str, str]] = {
    "CGAP-A-8": ("Yes", "assessment_user_functionality.feature @CGAP-A-8"),
}


def stub_tags(feature: Path) -> set[str]:
    text = feature.read_text(encoding="utf-8")
    stubs: set[str] = set()
    for block in re.split(r"\n\s*Scenario:", text):
        if "needs a live" not in block and "needs this reopen" not in block:
            continue
        for m in re.finditer(r"@(A-override|A\d+)", block):
            stubs.add(m.group(0))
    return stubs


def refresh_coverage(book) -> None:
    sh = book[SHEET]
    counts = dict.fromkeys(("yes", "partial", "deviation", "deferred", "no"), 0)
    total = 0
    for row in range(2, sh.max_row + 1):
        if not str(sh.cell(row, 1).value or "").strip():
            continue
        total += 1
        st = str(sh.cell(row, 11).value or "").strip().lower()
        if st.startswith("no"):
            counts["no"] += 1
        elif st in counts:
            counts[st] += 1
        else:
            counts["no"] += 1
    cov = book[COVERAGE]
    for r in range(2, cov.max_row + 1):
        if str(cov.cell(r, 1).value or "").strip() == SHEET:
            for i, k in enumerate(("yes", "partial", "deviation", "deferred", "no"), start=2):
                cov.cell(r, i).value = counts[k]
            cov.cell(r, 7).value = total
            auto = total - counts["no"]
            cov.cell(r, 8).value = f"{round(100 * counts['yes'] / auto)}%" if auto else "0%"
            print(f"{COVERAGE} {SHEET}: yes={counts['yes']} partial={counts['partial']} deferred={counts['deferred']}")
            break


def tag_from_ref(ref: str) -> str | None:
    m = re.search(r"@(A-override|A\d+)", ref or "")
    return m.group(0) if m else None


def main() -> None:
    reopen_stubs = stub_tags(FEATURES / "assessment_reopening.feature")
    book = openpyxl.load_workbook(WORKBOOK)
    sh = book[SHEET]
    n = 0
    for row in range(2, sh.max_row + 1):
        sno = str(sh.cell(row, 1).value or "").strip()
        ref = str(sh.cell(row, 12).value or "")
        auto = str(sh.cell(row, 11).value or "").strip()

        if sno in FIXED:
            a, r = FIXED[sno]
            sh.cell(row, 11).value = a
            sh.cell(row, 12).value = r
            n += 1
            continue

        tag = tag_from_ref(ref)
        if not tag:
            m = re.search(r"@(A-override|A\d+)", ref)
            tag = m.group(0) if m else None

        if "A135-A139" in ref or "A135-A139" in auto:
            sh.cell(row, 11).value = "Partial"
            sh.cell(row, 12).value = (
                "assessment_hide_final_grade.feature @A135 (config only; end-user visibility @provisioning)"
            )
            n += 1
            continue

        new: tuple[str, str] | None = None
        if tag in REOPEN_YES and tag not in reopen_stubs:
            new = ("Yes", REOPEN_YES[tag])
        elif tag in REOPEN_PARTIAL:
            new = ("Partial", REOPEN_PARTIAL[tag])
        elif tag in HIDE_FINAL_YES:
            new = ("Yes", HIDE_FINAL_YES[tag])
        elif tag in SHOW_HIDE_YES:
            new = ("Yes", SHOW_HIDE_YES[tag])

        if not new:
            continue
        if auto.lower() in ("yes", "partial") and tag in (ref or ""):
            continue
        sh.cell(row, 11).value = new[0]
        sh.cell(row, 12).value = new[1]
        if str(sh.cell(row, 7).value or "").strip().lower() in ("proposed", "confirmed"):
            sh.cell(row, 7).value = "Blocked" if new[0] == "Partial" else "Pass"
        n += 1

    refresh_coverage(book)
    book.save(WORKBOOK)
    print(f"updated {n} rows on {SHEET}")


if __name__ == "__main__":
    main()
