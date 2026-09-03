"""Merge Moodle_Guided_Labs into Practice_Project as one sheet.

Practice Project is the renamed Guided Project feature (the Practice_Project
sheet even carries a case for "all guided strings converted into practice
project"), so QA asked for a single sheet and a single release-plan row named
Practice_Project.

Automatic matching found NO exact and NO near-duplicate titles/summaries between
the sheets - the highest token overlap of any pair was 0.36 - because the guided
sheet describes the older flow while Practice_Project describes the renamed
feature plus its reports and attempts pages. The duplicates below were therefore
identified by reading both sheets: each entry is a guided row whose check already
exists in Practice_Project, with the practice row that covers it. Anything not
listed is carried over as a unique case.

Usage: python tools/merge_guided_into_practice.py
"""
from __future__ import annotations

import os
import shutil
import time
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
STAGING = Path("target/kiwi_guided_merge.xlsx")
SOURCE, TARGET = "Moodle_Guided_Labs", "Practice_Project"
COVERAGE = "Automation_Coverage"
LAST_COLUMN = 12  # A..L: the guided sheet's Automated + Automation_Ref included
CATEGORY_COL, AUTOMATED_COL, REF_COL = 5, 11, 12

# guided row -> (practice row already covering it, what the pair share)
DUPLICATES = {
    2: (2, "activity-type option on the activity settings page"),
    12: (131, "admin views the checkpoint details/results"),
    13: (126, "attempt start and completion dates shown to admin"),
    14: (125, "per-attempt checkpoint values are correct"),
    17: (11, "user stops the lab from the console"),
    26: (69, "the project report exists and can be filtered"),
    27: (137, "report data matches the UI/underlying data"),
    28: (114, "consolidated report can be downloaded"),
    30: (13, "instructions and access details on the console"),
    34: (25, "Start Lab after the lab was deleted re-creates it"),
    46: (42, "new attempt blocked during the between-attempts duration"),
    48: (20, "submitting an attempt runs the evaluation"),
    52: (11, "Stop Lab returns the user to the activity page"),
    53: (131, "checkpoint details modal renders the data"),
    54: (118, "admin drills into a user's attempts page"),
    55: (48, "admin overrides a user's attempts"),
    59: (16, "with vs without progress tracking activity types"),
}


def publish() -> None:
    openpyxl.load_workbook(STAGING)
    for attempt in range(10):
        try:
            os.replace(STAGING, WORKBOOK)
            return
        except PermissionError:
            if attempt == 9:
                raise
            time.sleep(1)


def add_missing_headers(target) -> None:
    headers = {11: "Automated", 12: "Automation_Ref"}
    donor = target.cell(1, 10)
    for column, label in headers.items():
        if not str(target.cell(1, column).value or "").strip():
            target.cell(1, column).value = label
            target.cell(1, column)._style = copy(donor._style)
            print(f"  added '{label}' header to {TARGET}")


def main() -> None:
    STAGING.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(WORKBOOK, STAGING)
    book = openpyxl.load_workbook(STAGING)
    source, target = book[SOURCE], book[TARGET]

    add_missing_headers(target)
    note = f"merged from {SOURCE} {date.today().isoformat()}"
    styles = [copy(target.cell(2, column)._style) for column in range(1, LAST_COLUMN + 1)]
    row = target.max_row + 1
    carried, skipped = [], []

    for source_row in range(2, source.max_row + 1):
        title = str(source.cell(source_row, 1).value or "").strip()
        summary = str(source.cell(source_row, 2).value or "").strip()
        if not title and not summary:
            continue
        if source_row in DUPLICATES:
            practice_row, reason = DUPLICATES[source_row]
            skipped.append((source_row, summary, practice_row, reason))
            continue

        for column in range(1, LAST_COLUMN + 1):
            target.cell(row, column).value = source.cell(source_row, column).value
            target.cell(row, column)._style = styles[column - 1]
        target.cell(row, CATEGORY_COL).value = TARGET
        existing = str(target.cell(row, REF_COL).value or "").strip()
        target.cell(row, REF_COL).value = f"{existing} | {note}" if existing else note
        carried.append((source_row, row, summary))
        row += 1

    del book[SOURCE]
    print(f"  deleted the {SOURCE} sheet")

    if COVERAGE in book.sheetnames:
        coverage = book[COVERAGE]
        for coverage_row in range(2, coverage.max_row + 1):
            if str(coverage.cell(coverage_row, 1).value or "").strip() == SOURCE:
                coverage.cell(coverage_row, 1).value = TARGET
                print(f"  relabelled {COVERAGE} row {coverage_row} to {TARGET} "
                      f"(its counts now need QA's numbers)")

    book.save(STAGING)
    publish()

    print(f"\nskipped {len(skipped)} duplicate cases:")
    for source_row, summary, practice_row, reason in skipped:
        print(f"  guided r{source_row:2d} -> practice r{practice_row:3d}  {reason}")
        print(f"      {summary[:110]}")
    total = sum(1 for check in range(2, target.max_row + 1)
                if str(target.cell(check, 1).value or "").strip()
                or str(target.cell(check, 2).value or "").strip())
    print(f"\ncarried over {len(carried)} unique cases into {TARGET} "
          f"(rows {carried[0][1]}-{carried[-1][1]}); the sheet now holds {total} cases")


if __name__ == "__main__":
    main()
