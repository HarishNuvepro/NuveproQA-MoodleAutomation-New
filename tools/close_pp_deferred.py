"""Mark remaining Practice Project deferred leftovers No. Leaves Deviations and existing Yes."""
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
ROWS = {
    13: "No — Credentials/PDF/help need a running configured lab (instructions/duration already in @PP-13)",
    56: "No — Not Attempted/Started override status needs an in-progress fixture (@PP-56 covers Completed)",
    177: "No — exhaust attempts needs a provisioned lab and attempt ceiling",
    178: "No — Complete is dormant (same product gap as PP-31/32)",
    179: "No — restart-in-gap needs a provisioned lab and cooldown",
    190: "No — override-card start needs a running lab plus override fixture",
    191: "No — last-attempt Stop/Delete needs a provisioned lab",
    192: "No — outside date window needs a dedicated dated activity (must not mutate GUIDED_CMID)",
}


def main() -> None:
    book = load_workbook(XLSX)
    sheet = book["Practice_Project"]
    for row, reason in ROWS.items():
        sheet.cell(row, 11).value = "No"
        sheet.cell(row, 12).value = reason
        sheet.cell(row, 7).value = "NA"
        print(f"  {row} No")
    book.save(XLSX)
    print("saved", XLSX)


if __name__ == "__main__":
    main()
