"""Promote Assessment rows whose Cucumber scenarios are implemented (medium batch 1c)."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Moodle_Assessment_Labs"
COVERAGE = "Automation_Coverage"
FEATURE_DIR = Path("src/test/resources/features")

BATCH1C_FEATURES = (
    "assessment_lab_actions.feature",
    "assessment_user_attempts_page.feature",
    "assessment_retake_cooldown.feature",
    "assessment_randomization_problem_statements.feature",
    "assessment_problem_statement_git_url.feature",
)

STUB_MARKERS = (
    "needs a live",
    "needs this reopen",
    "lifecycle case needs",
    "cooldown case needs",
)

# Provisioning scenarios: linked but expect skip without lab — Partial not Deferred
PROVISIONING_PARTIAL = {
    "@A108": "assessment_user_functionality.feature @A108 @provisioning",
    "@A109": "assessment_user_functionality.feature @A109 @provisioning",
    "@A8": "assessment_user_functionality.feature @A8 @provisioning",
    "@A114": "assessment_user_functionality.feature @A114 @provisioning",
    "@A113": "assessment_user_functionality.feature @A113 @provisioning",
    "@CGAP-A-9": "assessment_user_functionality.feature @CGAP-A-9 @provisioning",
}


def implemented_tags() -> dict[str, str]:
    """Map @tag -> 'feature @tag' for scenarios without stub Then lines."""
    out: dict[str, str] = {}
    for name in BATCH1C_FEATURES:
        path = FEATURE_DIR / name
        if not path.is_file():
            continue
        text = path.read_text(encoding="utf-8")
        chunks = re.split(r"(?=\n\s*@\S+\s*\n\s*Scenario:)", text)
        for chunk in chunks:
            if "Scenario:" not in chunk:
                continue
            if any(m in chunk for m in STUB_MARKERS):
                continue
            header = chunk.split("Scenario:", 1)[0]
            for m in re.finditer(r"@(A-override|A\d+[a-z]*)", header):
                tag = m.group(0)
                out[tag] = f"{name} {tag}"
    return out


def refresh_coverage(book) -> None:
    sh = book[SHEET]
    counts = dict.fromkeys(("yes", "partial", "deviation", "deferred", "no"), 0)
    total = 0
    for row in range(2, sh.max_row + 1):
        if not str(sh.cell(row, 1).value or "").strip():
            continue
        total += 1
        st = str(sh.cell(row, 11).value or "").strip().lower()
        if st.startswith("no"):
            counts["no"] += 1
        elif st in counts:
            counts[st] += 1
        else:
            counts["no"] += 1
    cov = book[COVERAGE]
    for r in range(2, cov.max_row + 1):
        if str(cov.cell(r, 1).value or "").strip() == SHEET:
            for i, k in enumerate(("yes", "partial", "deviation", "deferred", "no"), start=2):
                cov.cell(r, i).value = counts[k]
            cov.cell(r, 7).value = total
            auto = total - counts["no"]
            cov.cell(r, 8).value = f"{round(100 * counts['yes'] / auto)}%" if auto else "0%"
            print(f"{COVERAGE} {SHEET}: yes={counts['yes']} partial={counts['partial']} deferred={counts['deferred']}")
            break


def tags_in_text(text: str) -> list[str]:
    return re.findall(r"@(A-override|A\d+[a-z]*)", text or "")


def main() -> None:
    impl = implemented_tags()
    book = openpyxl.load_workbook(WORKBOOK)
    sh = book[SHEET]
    n = 0
    for row in range(2, sh.max_row + 1):
        if not str(sh.cell(row, 1).value or "").strip():
            continue
        ref = str(sh.cell(row, 12).value or "")
        auto = str(sh.cell(row, 11).value or "").strip().lower()
        sno = str(sh.cell(row, 1).value or "")

        prov_hit = next((t for t in PROVISIONING_PARTIAL if t in ref or t in sno), None)
        if prov_hit:
            if auto in ("yes", "partial") and prov_hit in ref:
                continue
            sh.cell(row, 11).value = "Partial"
            sh.cell(row, 12).value = PROVISIONING_PARTIAL[prov_hit]
            n += 1
            continue

        found = None
        for tag in tags_in_text(ref):
            if tag in impl:
                found = (tag, impl[tag])
                break
        if not found:
            m = re.search(r"\b(A\d+)\b", ref)
            if m:
                tag = "@" + m.group(1)
                if tag in impl:
                    found = (tag, impl[tag])
        if not found:
            continue
        tag, new_ref = found
        if auto in ("yes", "partial") and tag in ref:
            continue
        sh.cell(row, 11).value = "Yes"
        sh.cell(row, 12).value = new_ref
        if str(sh.cell(row, 7).value or "").strip().lower() in ("proposed", "confirmed"):
            sh.cell(row, 7).value = "Pass"
        n += 1

    refresh_coverage(book)
    book.save(WORKBOOK)
    print(f"implemented tags in batch1c features: {len(impl)}")
    print(f"updated {n} rows on {SHEET}")


if __name__ == "__main__":
    main()
