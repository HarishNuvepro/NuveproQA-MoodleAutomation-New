"""Sync CGAP-A-SB-* rows on Moodle_Assessment_Labs from assessment_sandbox_parity.feature."""
from __future__ import annotations

from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Moodle_Assessment_Labs"
COVERAGE = "Automation_Coverage"

PARTIAL: dict[str, str] = {
    "CGAP-A-SB-7": "assessment_sandbox_parity.feature @CGAP-A-SB-7",
    "CGAP-A-SB-8": "assessment_sandbox_parity.feature @CGAP-A-SB-8",
    "CGAP-A-SB-9": "assessment_sandbox_parity.feature @CGAP-A-SB-9",
    "CGAP-A-SB-10": "assessment_sandbox_parity.feature @CGAP-A-SB-10",
    "CGAP-A-SB-11": "assessment_sandbox_parity.feature @CGAP-A-SB-11",
    "CGAP-A-SB-13": "assessment_sandbox_parity.feature @CGAP-A-SB-13",
    "CGAP-A-SB-14": "assessment_sandbox_parity.feature @CGAP-A-SB-14",
    "CGAP-A-SB-15": "assessment_sandbox_parity.feature @CGAP-A-SB-15",
    "CGAP-A-SB-16": "assessment_sandbox_parity.feature @CGAP-A-SB-16",
    "CGAP-A-SB-17": "assessment_sandbox_parity.feature @CGAP-A-SB-17",
    "CGAP-A-SB-18": "assessment_sandbox_parity.feature @CGAP-A-SB-18",
    "CGAP-A-SB-19": "assessment_sandbox_parity.feature @CGAP-A-SB-19",
    "CGAP-A-SB-20": "assessment_sandbox_parity.feature @CGAP-A-SB-20",
    "CGAP-A-SB-21": "assessment_sandbox_parity.feature @CGAP-A-SB-21",
    "CGAP-A-SB-23": "assessment_sandbox_parity.feature @CGAP-A-SB-23",
    "CGAP-A-SB-26": "assessment_sandbox_parity.feature @CGAP-A-SB-26",
    "CGAP-A-SB-27": "assessment_sandbox_parity.feature @CGAP-A-SB-27",
    "CGAP-A-SB-28": "assessment_sandbox_parity.feature @CGAP-A-SB-28",
}

PARTIAL["CGAP-A-SB-12"] = "assessment_user_functionality.feature @A-user-22"

DEFERRED: dict[str, str] = {
    "CGAP-A-SB-22": "Deferred — bulk Create from No labs (port @SB-manage-12)",
    "CGAP-A-SB-24": "Deferred — catalog/unlimited assessment fixture",
    "CGAP-A-SB-25": "Deferred — invalid-template assessment fixture",
}


def refresh_coverage(book) -> None:
    sh = book[SHEET]
    counts = dict.fromkeys(("yes", "partial", "deviation", "deferred", "no"), 0)
    total = 0
    for row in range(2, sh.max_row + 1):
        if not str(sh.cell(row, 1).value or "").strip():
            continue
        total += 1
        st = str(sh.cell(row, 11).value or "").strip().lower()
        if st.startswith("no"):
            counts["no"] += 1
        elif st in counts:
            counts[st] += 1
        else:
            counts["no"] += 1
    cov = book[COVERAGE]
    for r in range(2, cov.max_row + 1):
        if str(cov.cell(r, 1).value or "").strip() == SHEET:
            cov.cell(r, 2).value = counts["yes"]
            cov.cell(r, 3).value = counts["partial"]
            cov.cell(r, 4).value = counts["deviation"]
            cov.cell(r, 5).value = counts["deferred"]
            cov.cell(r, 6).value = counts["no"]
            cov.cell(r, 7).value = total
            auto = total - counts["no"]
            cov.cell(r, 8).value = f"{round(100 * counts['yes'] / auto)}%" if auto else "0%"
            break


def apply() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sh = book[SHEET]
    updated = 0
    for row in range(2, sh.max_row + 1):
        sno = str(sh.cell(row, 1).value or "").strip()
        if sno in PARTIAL:
            sh.cell(row, 11).value = "Partial"
            sh.cell(row, 12).value = PARTIAL[sno]
            sh.cell(row, 7).value = "Blocked"
            updated += 1
        elif sno in DEFERRED:
            sh.cell(row, 11).value = "Deferred"
            sh.cell(row, 12).value = DEFERRED[sno]
            sh.cell(row, 7).value = "NA"
            updated += 1
    refresh_coverage(book)
    book.save(WORKBOOK)
    print(f"updated {updated} CGAP-A-SB rows on {SHEET}; saved {WORKBOOK}")


if __name__ == "__main__":
    apply()
