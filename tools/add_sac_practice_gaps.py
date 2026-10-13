"""Append CGAP-SAC-PP-* Practice Project sequential completion rows to Kiwi."""
from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Sequential_Activity_Completion"
FEATURE = "sequential_completion.feature"
TODAY = date.today().isoformat()

GAPS = [
    (
        "CGAP-SAC-PP-1",
        "Practice Project mark-as-done restriction — negative",
        "Validate Activity 2 remains restricted when Activity 1 is not marked done (Practice Project, manual completion).",
        "Practice Project Activity 1 uses manual mark-as-done; Activity 2 restricted until Activity 1 complete.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-1",
    ),
    (
        "CGAP-SAC-PP-2",
        "Practice Project mark-as-done — incomplete status",
        "Validate Activity 1 shows incomplete on the course page before mark-as-done (Practice Project).",
        "Manual mark-as-done on Activity 1.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-2",
    ),
    (
        "CGAP-SAC-PP-3",
        "Practice Project restrict message",
        "Restriction message names Activity 1 when Activity 2 is gated on Activity 1 completion (Practice Project).",
        "Activity 2 restricted until Activity 1 complete.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-3",
    ),
    (
        "CGAP-SAC-PP-4",
        "Practice Project mark-as-done — completed status",
        "Activity 1 shows completed after the student marks as done (Practice Project).",
        "Manual mark-as-done on Activity 1.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-4",
    ),
    (
        "CGAP-SAC-PP-5",
        "Practice Project mark-as-done — unlock Activity 2",
        "Activity 2 becomes accessible once Activity 1 is marked done (Practice Project).",
        "Manual mark-as-done on Activity 1; Activity 2 restricted until Activity 1 complete.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-5",
    ),
    (
        "CGAP-SAC-PP-6",
        "Practice Project mark-as-done — persistence",
        "Activity 2 remains accessible after logout/login once Activity 1 is marked done (Practice Project).",
        "Activity 2 unlocked after Activity 1 mark-as-done.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-6",
    ),
    (
        "CGAP-SAC-PP-7",
        "Practice Project three-activity chain",
        "Activity 3 stays locked until both Activity 1 and Activity 2 are marked done (Practice Project).",
        "Three Practice Project activities with sequential restrict-access.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-7",
    ),
    (
        "CGAP-SAC-PP-8",
        "Practice Project activity completion — negative",
        "Activity 2 stays restricted until the student completes Activity 1 (Practice Project, no grade rule).",
        "Activity 1 completion = student must complete the activity; Activity 2 restricted until Activity 1 complete.",
        "High",
        f"{FEATURE} @CGAP-SAC-PP-8",
    ),
    (
        "CGAP-SAC-PP-9",
        "Practice Project activity completion — unlock",
        "Activity 2 unlocks after the student completes Activity 1 (submit Practice Project attempt).",
        "Student must complete Activity 1; CloudLabs lab provisioning.",
        "High",
        f"{FEATURE} @CGAP-SAC-PP-9",
    ),
    (
        "CGAP-SAC-PP-10",
        "Practice Project activity completion — incomplete",
        "Activity 1 shows incomplete before the student completes the Practice Project.",
        "Student must complete Activity 1.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-10",
    ),
    (
        "CGAP-SAC-PP-11",
        "Practice Project activity completion — view only",
        "Activity 2 stays locked when the student opens Activity 1 but does not complete it.",
        "Student must complete Activity 1.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-11",
    ),
    (
        "CGAP-SAC-PP-12",
        "Practice Project activity completion — persistence",
        "Activity 2 remains accessible after re-login once Activity 1 was completed.",
        "Activity 1 completed via Practice Project submit.",
        "Medium",
        f"{FEATURE} @CGAP-SAC-PP-12",
    ),
]


def next_row(sheet) -> int:
    for row in range(sheet.max_row, 1, -1):
        if str(sheet.cell(row, 1).value or "").strip():
            return row + 1
    return 2


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sheet = book[SHEET]
    existing = {
        str(sheet.cell(r, 1).value or "").strip()
        for r in range(2, sheet.max_row + 1)
    }
    row = next_row(sheet)
    added = 0
    for gap_id, title, summary, pre, priority, ref in GAPS:
        if gap_id in existing:
            print(f"skip duplicate {gap_id}")
            continue
        sheet.cell(row, 1).value = gap_id
        sheet.cell(row, 2).value = title
        sheet.cell(row, 3).value = pre or summary
        sheet.cell(row, 4).value = "Moodle"
        sheet.cell(row, 5).value = "Sequential_Activity_Completion"
        sheet.cell(row, 6).value = priority
        sheet.cell(row, 7).value = "Pass"
        sheet.cell(row, 8).value = ""
        sheet.cell(row, 9).value = ""
        sheet.cell(row, 10).value = "Practice Project"
        sheet.cell(row, 11).value = "Yes"
        sheet.cell(row, 12).value = ref
        sheet.cell(row, 13).value = f"Automation gap (Practice Project sequential) {TODAY}"
        row += 1
        added += 1
    book.save(WORKBOOK)
    print(f"added {added} rows to {SHEET} in {WORKBOOK}")


if __name__ == "__main__":
    main()
