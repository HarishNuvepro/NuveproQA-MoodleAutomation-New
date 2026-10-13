"""Set Case_Status (Pass/Fail) on easy batch rows from merged Cucumber JSON."""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "tools") not in sys.path:
    sys.path.insert(0, str(ROOT / "tools"))

from traceability import norm_tag, pick_result, read_execution, tags_in  # noqa: E402

BATCH = ROOT / "reports" / "easy_automation_batches.csv"
XLSX = ROOT / "docs" / "MoodleRegression_Testing_Kiwi new.xlsx"
STATUS_COL = 7  # openpyxl 1-based (Case_Status)


def primary_tag(case_id: str, automation_ref: str) -> str | None:
    if case_id.startswith("CGAP-"):
        return case_id
    refs = tags_in(automation_ref or "")
    return refs[0] if refs else None


def main() -> None:
    outcomes, used = read_execution()
    if not used:
        print("No cucumber JSON — skip Case_Status updates")
        return
    rows = list(csv.DictReader(BATCH.open(encoding="utf-8")))
    wb = openpyxl.load_workbook(XLSX)
    updated = 0
    for row in rows:
        area = row["area"]
        cid = row["case_id"]
        ws = wb[area]
        for r in range(2, ws.max_row + 1):
            sno = str(ws.cell(r, 1).value or "").strip()
            title = str(ws.cell(r, 2).value or "").strip()
            if cid.startswith("CGAP-"):
                if sno != cid:
                    continue
            elif sno != cid and title != cid:
                continue
            ref = str(ws.cell(r, 12).value or "")
            tag = primary_tag(cid, ref)
            lookup_tags = tags_in(ref)
            if tag:
                lookup_tags = [tag] + [t for t in lookup_tags if t.upper() != tag.upper()]
            if not lookup_tags:
                continue
            oc = pick_result(lookup_tags, outcomes)
            if oc == "passed":
                ws.cell(r, STATUS_COL).value = "Pass"
                updated += 1
            elif oc == "failed":
                ws.cell(r, STATUS_COL).value = "Fail"
                updated += 1
            break
    wb.save(XLSX)
    print(f"Updated Case_Status on {updated} easy-batch rows (sources: {', '.join(used)})")


if __name__ == "__main__":
    main()
