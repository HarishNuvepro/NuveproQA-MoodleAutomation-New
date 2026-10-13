"""Mark Practice Project Moodle configurable-report cases as not automatable (Automated=No)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Practice_Project"
COVERAGE = "Automation_Coverage"

REASON = (
    "No — Not automatable: Moodle site configurable reports (detailed/summary Report Builder); "
    "deferred per QA — no stable report URLs on automation QA"
)

# Excel row numbers on Practice_Project (same set as add_practice_project_gaps CONVERTED_NO reports).
REPORT_ROWS = sorted({
    69, 70, 71, 72,
    *range(73, 93),
    *range(94, 113),
    137, 138, 142,
})


def refresh_coverage_row(book) -> None:
    sheet = book[SHEET]
    states = ("yes", "partial", "deviation", "deferred", "no")
    counts = dict.fromkeys(states, 0)
    total = 0
    for row in range(2, sheet.max_row + 1):
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        total += 1
        state = str(sheet.cell(row, 11).value or "").strip().lower()
        if state in counts:
            counts[state] += 1
    if COVERAGE not in book.sheetnames:
        return
    coverage = book[COVERAGE]
    target = None
    for row in range(2, coverage.max_row + 1):
        if str(coverage.cell(row, 1).value or "").strip() == SHEET:
            target = row
            break
    if target is None:
        return
    values = [counts[s] for s in states] + [total]
    for offset, value in enumerate(values, start=2):
        coverage.cell(target, offset).value = value
    coverage.cell(target, 8).value = f"{round(100 * counts['yes'] / total)}%" if total else "0%"
    print(f"{COVERAGE} {SHEET}: {values} ({coverage.cell(target, 8).value})")


def mark_by_ref(book) -> int:
    """Catch report rows missed by row index (ref mentions practice_project_reports)."""
    sheet = book[SHEET]
    extra = 0
    for row in range(2, sheet.max_row + 1):
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        ref = str(sheet.cell(row, 12).value or "")
        if "practice_project_reports" not in ref.lower() and "configurable-report" not in ref.lower():
            continue
        auto = str(sheet.cell(row, 11).value or "").strip().lower()
        if auto == "yes":
            continue
        sheet.cell(row, 11).value = "No"
        sheet.cell(row, 12).value = REASON
        sheet.cell(row, 7).value = "NA"
        extra += 1
    return extra


def main() -> None:
    book = load_workbook(WORKBOOK)
    sheet = book[SHEET]
    marked = 0
    for row in REPORT_ROWS:
        if row > sheet.max_row:
            continue
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        auto = str(sheet.cell(row, 11).value or "").strip().lower()
        if auto == "yes":
            print(f"  skip row {row} (already Yes)")
            continue
        sheet.cell(row, 11).value = "No"
        sheet.cell(row, 12).value = REASON
        sheet.cell(row, 7).value = "NA"
        marked += 1
    extra = mark_by_ref(book)
    refresh_coverage_row(book)
    book.save(WORKBOOK)
    print(f"marked {marked} report rows by index + {extra} by ref; saved {WORKBOOK}")


if __name__ == "__main__":
    main()
