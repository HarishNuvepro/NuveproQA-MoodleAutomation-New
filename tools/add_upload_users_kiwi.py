"""Create the derived Upload_Users Kiwi sheet and dashboard rows in place. No backups."""
from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "MoodleRegression_Testing_Kiwi new.xlsx"
SHEET = "Upload_Users"

# id, priority, summary, expected, automation status, automation ref
CASES = [
    ("UU-01", "High", "Admin can open Site administration > Users > Accounts > Upload users",
     "Upload form shows the file, encoding, delimiter and preview-row controls.", "No", "@UU-S1"),
    ("UU-02", "High", "The example Upload users CSV can be downloaded",
     "Downloaded CSV contains the required username, firstname, lastname and email headers.", "No", "@UU-S1"),
    ("UU-03", "High", "A valid three-user CSV creates all users",
     "Results report three created users; all are listed and one can log in.", "No", "@UU-S1"),
    ("UU-04", "High", "Preview shows parsed CSV rows before upload is committed",
     "Preview lists the CSV usernames and no account exists before Upload users is clicked.", "No", "@UU-S2"),
    ("UU-05", "High", "Cancel from preview does not create users",
     "The preview is left and none of its usernames exists.", "No", "@UU-S2"),
    ("UU-06", "High", "No file and an empty file are rejected",
     "A validation error is shown and no account is created.", "No", "@UU-S3"),
    ("UU-07", "High", "A CSV missing a required username or email header is rejected",
     "Preview/upload reports the missing field and creates no account.", "No", "@UU-S4"),
    ("UU-08", "High", "A mixed valid and malformed CSV reports row-level errors without aborting",
     "The valid row is created; missing email, invalid email and invalid username rows report errors.", "No", "@UU-S5"),
    ("UU-09", "High", "Add new only skips an existing username",
     "Existing account is unchanged while new rows are created.", "No", "@UU-S6"),
    ("UU-10", "High", "Add new and update existing overrides an existing user from file",
     "Existing details are updated and a new CSV user is created.", "No", "@UU-S7"),
    ("UU-11", "High", "Update existing only updates matches and ignores unknown usernames",
     "Existing account is updated and the unknown username is not created.", "No", "@UU-S8"),
    ("UU-12", "Medium", "Add all appends a number when the username already exists",
     "A second account is created with a numeric suffix and the original is unchanged.", "No", "@UU-S17"),
    ("UU-13", "High", "Prevent email duplicates rejects a new username with an existing email",
     "A row error is reported and the duplicate-email account is not created.", "No", "@UU-S9"),
    ("UU-14", "High", "A password supplied by CSV is usable",
     "The uploaded user can log in with the supplied password.", "No", "@UU-S1"),
    ("UU-15", "High", "A weak CSV password is rejected when password is required in file",
     "The row reports a password-policy error and no usable account is created.", "No", "@UU-S10"),
    ("UU-16", "High", "suspended=1 suspends an existing user when suspending is allowed",
     "User is flagged suspended and cannot log in.", "No", "@UU-S11"),
    ("UU-17", "High", "suspended=0 reactivates a suspended existing user",
     "User is active and can log in again.", "No", "@UU-S11"),
    ("UU-18", "High", "deleted=1 deletes an existing user when deletes are allowed",
     "User is absent from the list and cannot log in.", "No", "@UU-S12"),
    ("UU-19", "Medium", "deleted=1 is ignored when deletes are not allowed",
     "User remains listed and can still log in.", "No", "@UU-S12"),
    ("UU-20", "Medium", "oldusername renames an existing user when renames are allowed",
     "New username is listed and can be used to log in.", "No", "@UU-S13"),
    ("UU-21", "Medium", "oldusername is ignored when renames are not allowed",
     "Existing username remains unchanged.", "No", "@UU-S13"),
    ("UU-22", "High", "CSV cannot delete a site administrator account",
     "Administrator remains active after deleted=1 with Allow deletes enabled.", "No", "@UU-S18"),
    ("UU-23", "High", "course1 enrols the uploaded user in an existing course shortname",
     "User opens the course page as an enrolled student.", "No", "@UU-S14"),
    ("UU-24", "Medium", "An unknown course1 shortname reports a warning without creating a course",
     "User is created but no bogus course or enrolment is created.", "No", "@UU-S14"),
    ("UU-25", "Medium", "cohort1 adds the uploaded user to an existing cohort ID",
     "User appears as a member of the specified cohort.", "No", "@UU-S15"),
    ("UU-26", "Medium", "Preview default City is applied when the CSV omits city",
     "Created profile contains the configured default city.", "No", "@UU-S15"),
    ("UU-27", "Medium", "A delimiter mismatch is rejected",
     "CSV cannot be parsed with the selected delimiter and creates no account.", "No", "@UU-S16"),
    ("UU-28", "High", "A student cannot access Upload users",
     "Direct URL is denied and the Upload users form is not shown.", "No", "@UU-S16"),
    ("UU-29", "Medium", "Continue from results returns to Upload users",
     "Upload users initial form is displayed again.", "No", "@UU-S1"),
    ("UU-30", "Medium", "Optional profile fields in CSV persist",
     "City, country and department values appear on the created profile.", "No", "@UU-S15"),
    ("UU-31", "Low", "A non-UTF-8 file is handled according to the selected encoding",
     "Preview decodes names correctly or reports an encoding error.", "Deferred", "Deferred encoding variant"),
    ("UU-32", "Low", "Preview rows setting limits the displayed preview",
     "Preview displays no more than the selected number of records.", "Deferred", "Deferred cosmetic preview variant"),
    ("UU-33", "Medium", "Force password change All is enforced for uploaded users",
     "Uploaded user must change the password at first login.", "Deferred", "Deferred; overlaps Moodle_users password flow"),
    ("UU-34", "Medium", "Existing user password Update replaces the password from CSV",
     "Old password fails and CSV password succeeds.", "Deferred", "Deferred extra login state flow"),
    ("UU-35", "Medium", "Fill in missing updates only empty existing user details",
     "Existing non-empty fields remain and empty fields receive CSV/default values.", "Deferred", "Deferred fine-grained overlay"),
    ("UU-36", "Medium", "Standardise usernames converts uploaded usernames",
     "Username is normalized according to the selected setting.", "Deferred", "Deferred username policy variant"),
    ("UU-37", "Medium", "Select for bulk user actions forwards selected uploaded users",
     "Bulk user actions page contains the selected created/updated users.", "Deferred", "Deferred different admin tool"),
    ("UU-38", "Medium", "group1 assigns the user and creates a missing course group",
     "User belongs to the specified group after upload.", "Deferred", "Deferred course/group fixture"),
    ("UU-39", "Medium", "role1 or type1 assigns the configured course role",
     "Uploaded enrolment has the specified role.", "Deferred", "Deferred role variants"),
    ("UU-40", "Medium", "enrolperiod1 and enrolstatus1 control course enrolment",
     "Enrolment duration/status matches the CSV.", "Deferred", "Deferred time/course suspension"),
    ("UU-41", "Medium", "sysrole1 assigns and removes a system role",
     "User system-role assignment matches the CSV.", "Deferred", "Deferred; overlaps Moodle_users system role"),
    ("UU-42", "Medium", "profile_field_* imports a configured custom profile field",
     "Custom profile value persists.", "Deferred", "Deferred; needs site custom field"),
    ("UU-43", "Medium", "Create password if needed generates and emails a password",
     "User receives usable generated credentials.", "Deferred", "Deferred; requires mail capture"),
    ("UU-44", "Low", "Default-value username templates expand user fields",
     "Template substitutions produce the expected values.", "Deferred", "Deferred low-priority template matrix"),
    ("UU-45", "Low", "Values over database field limits report a row error",
     "Oversized record is not created and error is clear.", "Deferred", "Deferred low-priority boundaries"),
    ("UU-46", "Low", "Trailing spaces and encoded commas are handled",
     "Whitespace/errors and &#44 decoding follow Moodle behaviour.", "Deferred", "Deferred low-priority format edge"),
    ("UU-47", "Medium", "auth=nologin creates a user that cannot log in",
     "Account is listed with No login authentication and login is rejected.", "Deferred", "Deferred auth-plugin variant"),
    ("UU-48", "Medium", "Duplicate emails can be allowed when the site setting permits it",
     "Two usernames can share an email only when site and upload settings allow it.", "Deferred", "Deferred site authentication mutation"),
]


def copy_sheet_style(source, target) -> None:
    for c in range(1, 13):
        src = source.cell(1, c)
        dst = target.cell(1, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        target.column_dimensions[dst.column_letter].width = source.column_dimensions[src.column_letter].width
    target.freeze_panes = source.freeze_panes
    target.auto_filter.ref = f"A1:L{len(CASES) + 1}"


def add_automation_coverage(wb) -> None:
    ws = wb["Automation_Coverage"]
    existing = next((r for r in range(2, ws.max_row + 1)
                     if str(ws.cell(r, 1).value or "").strip() == SHEET), None)
    if existing is None:
        anchor = next((r for r in range(2, ws.max_row + 1)
                       if str(ws.cell(r, 1).value or "").strip() == "Moodle_users"), 1)
        existing = anchor + 1
        ws.insert_rows(existing)
    values = [SHEET, 0, 0, 0, 18, 30, 48, "0%"]
    for c, value in enumerate(values, 1):
        ws.cell(existing, c).value = value
    total_row = next((r for r in range(2, ws.max_row + 1)
                      if str(ws.cell(r, 1).value or "").startswith("TOTAL")), None)
    if total_row:
        totals = [sum(int(ws.cell(r, c).value or 0) for r in range(2, total_row)
                      if ws.cell(r, 1).value)
                  for c in range(2, 7)]
        grand = sum(totals)
        for c, value in enumerate(totals, 2):
            ws.cell(total_row, c).value = value
        ws.cell(total_row, 7).value = grand
        ws.cell(total_row, 8).value = f"{round(100 * totals[0] / grand)}%" if grand else "0%"


def add_regression_dashboard(wb) -> None:
    ws = wb["Regression_Dashboard"]
    # Remove stale row if this script is rerun.
    stale = next((r for r in range(5, ws.max_row + 1)
                  if str(ws.cell(r, 2).value or "").strip() == SHEET), None)
    if stale:
        ws.delete_rows(stale)
    anchor = next((r for r in range(5, ws.max_row + 1)
                   if str(ws.cell(r, 2).value or "").strip() == "Moodle_users"), 5)
    row = anchor + 1
    ws.insert_rows(row)
    ws.cell(row, 1).value = 38
    ws.cell(row, 2).value = SHEET
    ws.cell(row, 3).value = f"=COUNTA({SHEET}!A2:A2000)"
    ws.cell(row, 5).value = (
        f'=COUNTIFS({SHEET}!K2:K2000,"Yes",{SHEET}!G2:G2000,"Pass")'
        f'+COUNTIFS({SHEET}!K2:K2000,"Yes",{SHEET}!G2:G2000,"Fail")'
        f'+COUNTIFS({SHEET}!K2:K2000,"Yes",{SHEET}!G2:G2000,"Failed")'
        f'+COUNTIFS({SHEET}!K2:K2000,"Yes",{SHEET}!G2:G2000,"Blocked")')
    ws.cell(row, 6).value = f'=COUNTIF({SHEET}!K2:K2000,"No")'
    ws.cell(row, 7).value = f"=MAX(0,C{row}-E{row}-F{row})"
    ws.cell(row, 8).value = f'=IF((C{row}-F{row})<=0,0,E{row}/(C{row}-F{row}))'
    # All row-relative dashboard formulas and TOTAL must follow the insertion.
    total = None
    for r in range(5, ws.max_row + 1):
        if str(ws.cell(r, 2).value or "").strip().upper() == "TOTAL":
            total = r
            break
        if r != row and ws.cell(r, 2).value:
            ws.cell(r, 7).value = f"=MAX(0,C{r}-E{r}-F{r})"
            ws.cell(r, 8).value = f'=IF((C{r}-F{r})<=0,0,E{r}/(C{r}-F{r}))'
    if total:
        for c, col in ((3, "C"), (5, "E"), (6, "F"), (7, "G")):
            ws.cell(total, c).value = f"=SUM({col}5:{col}{total - 1})"
        ws.cell(total, 8).value = f'=IF((C{total}-F{total})<=0,0,E{total}/(C{total}-F{total}))'


def main() -> None:
    wb = load_workbook(XLSX)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    source = wb["Moodle_users"]
    index = wb.sheetnames.index("Moodle_users") + 1
    ws = wb.create_sheet(SHEET, index)
    headers = ["Title", "Summary", "Preconditions", "Product", "Category", "Priority",
               "Case_Status", "Steps", "Expected", "Regression_Tier", "Automated", "Automation_Ref"]
    for c, value in enumerate(headers, 1):
        ws.cell(1, c).value = value
    copy_sheet_style(source, ws)
    wrap = Alignment(wrap_text=True, vertical="top")
    for r, (case_id, priority, summary, expected, automated, ref) in enumerate(CASES, 2):
        values = [
            case_id, summary,
            "Site administrator account; Moodle Upload users page available",
            "Moodle", "Upload Users", priority, "PROPOSED",
            "Derived from Moodle 4.5 Upload users control inventory, CSV state model, and role/input matrices.",
            expected, "Core" if priority == "High" else "Extended",
            automated,
            (f"upload_users.feature {ref} | derived 2026-08-26"
             if automated == "No" else ref),
        ]
        for c, value in enumerate(values, 1):
            ws.cell(r, c).value = value
            ws.cell(r, c).alignment = wrap
    add_automation_coverage(wb)
    add_regression_dashboard(wb)
    tmp = XLSX.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    wb.close()
    tmp.replace(XLSX)
    print(f"Saved {XLSX}: {SHEET} with {len(CASES)} cases (30 convert, 18 deferred)")


if __name__ == "__main__":
    main()
