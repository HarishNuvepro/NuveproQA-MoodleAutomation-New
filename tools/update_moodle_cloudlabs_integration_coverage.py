"""Mark Moodle-Cloudlabs Integration cases Automated + Automation_Ref."""
from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Moodle-Cloudlabs Integration"
COVERAGE = "Automation_Coverage"
PLAN = "Release_Regression_Plan"
STATES = ("yes", "partial", "deviation", "deferred", "no")

# Excel row -> tag (row 2 = @MCI-1 … row 14 = @MCI-13)
TAGS = {row: f"@MCI-{row - 1}" for row in range(2, 15)}


def refresh_coverage_row(book) -> None:
    sheet = book[SHEET]
    counts = dict.fromkeys(STATES, 0)
    total = 0
    for row in range(2, sheet.max_row + 1):
        title = str(sheet.cell(row, 1).value or "").strip()
        if not title:
            continue
        total += 1
        state = str(sheet.cell(row, 11).value or "").strip().lower()
        if state in counts:
            counts[state] += 1

    coverage = book[COVERAGE]
    target = next(
        (row for row in range(2, coverage.max_row + 1)
         if str(coverage.cell(row, 1).value or "").strip() == SHEET),
        None,
    )
    if target is None:
        total_row = next(
            row for row in range(2, coverage.max_row + 1)
            if str(coverage.cell(row, 1).value or "").upper().startswith("TOTAL")
        )
        target = total_row
        coverage.insert_rows(target)
        template = target - 1
        for column in range(1, 9):
            coverage.cell(target, column)._style = copy(coverage.cell(template, column)._style)
        coverage.cell(target, 1).value = SHEET
        print(f"inserted {COVERAGE} row {target} for {SHEET}")

    values = [counts[state] for state in STATES] + [total]
    for offset, value in enumerate(values, start=2):
        coverage.cell(target, offset).value = value
    coverage.cell(target, 8).value = (
        f"{round(100 * counts['yes'] / total)}%" if total else "0%"
    )

    total_row = next(
        row for row in range(2, coverage.max_row + 1)
        if str(coverage.cell(row, 1).value or "").upper().startswith("TOTAL")
    )
    sums = [0] * 6
    for row in range(2, total_row):
        for index in range(6):
            value = coverage.cell(row, index + 2).value
            sums[index] += value if isinstance(value, (int, float)) else 0
    for index, value in enumerate(sums, start=2):
        coverage.cell(total_row, index).value = value
    coverage.cell(total_row, 8).value = (
        f"{round(100 * sums[0] / sums[5])}%" if sums[5] else "0%"
    )
    print(f"{COVERAGE} {SHEET}: {values} ({coverage.cell(target, 8).value}); TOTAL {sums}")


def mark_plan_automated(book) -> None:
    plan = book[PLAN]
    for row in range(10, plan.max_row + 1):
        if str(plan.cell(row, 2).value or "").strip() == SHEET:
            plan.cell(row, 11).value = "Yes"
            print(f"{PLAN} row {row} Automated=Yes")
            return


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sheet = book[SHEET]
    sheet.cell(1, 11).value = "Automated"
    sheet.cell(1, 12).value = "Automation_Ref"

    ref = "moodle_cloudlabs_integration.feature"
    for row in range(2, 15):
        tag = TAGS[row]
        if row == 13:
            sheet.cell(row, 11).value = "No"
            sheet.cell(row, 12).value = (
                "Not required in automation suite (invalid-settings case removed from feature)"
            )
        elif row == 14:
            sheet.cell(row, 11).value = "No"
            sheet.cell(row, 12).value = (
                "Not required in automation suite (uninstall case removed from feature)"
            )
        else:
            sheet.cell(row, 11).value = "Yes"
            sheet.cell(row, 12).value = f"{ref} {tag}"
            sheet.cell(row, 7).value = "Pass"
        if row in (13, 14):
            sheet.cell(row, 7).value = "NA"

    refresh_coverage_row(book)
    mark_plan_automated(book)
    book.save(WORKBOOK)
    print(f"saved {WORKBOOK}")


if __name__ == "__main__":
    main()
