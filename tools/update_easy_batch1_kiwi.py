"""Mark batch-1 easy cases Yes/Partial with Automation_Ref after scenarios exist."""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BATCH = ROOT / "reports" / "easy_automation_batches.csv"
XLSX = ROOT / "docs" / "MoodleRegression_Testing_Kiwi new.xlsx"

# case_id (S.No) -> (Automated, Automation_Ref)
REF = {
    "CGAP-GR-1": ("Yes", "easy_automation_batch1.feature @CGAP-GR-1"),
    "CGAP-GR-2": ("Yes", "easy_automation_batch1.feature @CGAP-GR-2"),
    "CGAP-GR-3": ("Partial", "easy_automation_batch1.feature @CGAP-GR-3 — groups index smoke"),
    "CGAP-GR-4": ("Partial", "easy_automation_batch1.feature @CGAP-GR-4 — auto-create control smoke"),
    "CGAP-GR-5": ("Yes", "easy_automation_batch1.feature @CGAP-GR-5"),
    "CGAP-GR-6": ("Yes", "easy_automation_batch1.feature @CGAP-GR-6"),
    "CGAP-GR-7": ("Yes", "easy_automation_batch1.feature @CGAP-GR-7"),
    "CGAP-GR-8": ("Partial", "easy_automation_batch1.feature @CGAP-GR-8 — cohort-sync smoke"),
    "CGAP-GR-9": ("Partial", "easy_automation_batch1.feature @CGAP-GR-9 — import control smoke"),
    "CGAP-GR-10": ("Yes", "easy_automation_batch1.feature @CGAP-GR-10"),
    "CGAP-GR-11": ("Partial", "easy_automation_batch1.feature @CGAP-GR-11 — cohort-sync smoke"),
    "CGAP-R-1": ("Yes", "easy_automation_batch1.feature @CGAP-R-1"),
    "CGAP-R-2": ("Yes", "easy_automation_batch1.feature @CGAP-R-2"),
    "CGAP-R-3": ("Yes", "easy_automation_batch1.feature @CGAP-R-3"),
    "CGAP-R-6": ("Yes", "easy_automation_batch1.feature @CGAP-R-6"),
    "CGAP-R-7": ("Yes", "easy_automation_batch1.feature @CGAP-R-7"),
    "CGAP-R-10": ("Partial", "easy_automation_batch1.feature @CGAP-R-10 @skip"),
    "CGAP-R-12": ("Yes", "easy_automation_batch1.feature @CGAP-R-12"),
    "CGAP-T-1": ("Partial", "easy_automation_batch1.feature @CGAP-T-1 @skip"),
    "CGAP-T-2": ("Yes", "easy_automation_batch1.feature @CGAP-T-2"),
    "CGAP-T-4": ("Yes", "easy_automation_batch1.feature @CGAP-T-4"),
    "CGAP-T-5": ("Yes", "easy_automation_batch1.feature @CGAP-T-5"),
    "CGAP-T-7": ("Yes", "easy_automation_batch1.feature @CGAP-T-7"),
    "CGAP-T-10": ("Yes", "easy_automation_batch1.feature @CGAP-T-10"),
    "CGAP-ML-2": ("Partial", "easy_automation_batch1.feature @CGAP-ML-2 — needs -Pslow lab seed"),
    "CGAP-CIG-1": ("Yes", "easy_automation_batch1.feature @CGAP-CIG-1"),
}


def main() -> None:
    batch1 = []
    with BATCH.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row.get("batch") == "1":
                batch1.append(row)

    wb = openpyxl.load_workbook(XLSX)
    updated = 0
    for row in batch1:
        cid = row["case_id"]
        area = row["area"]
        if cid.startswith("CGAP-") and cid in REF:
            auto, ref = REF[cid]
        elif row.get("run_tag"):
            tag = row["run_tag"]
            feat = row.get("feature") or "labtemplates.feature"
            auto, ref = ("Partial", f"{feat} @{tag} @easybatch1")
        else:
            continue
        ws = wb[area]
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "").strip() != cid:
                continue
            ws.cell(r, 11).value = auto
            ws.cell(r, 12).value = ref
            updated += 1
            break
    wb.save(XLSX)
    print(f"Updated {updated} Kiwi rows for easy batch 1")


if __name__ == "__main__":
    main()
