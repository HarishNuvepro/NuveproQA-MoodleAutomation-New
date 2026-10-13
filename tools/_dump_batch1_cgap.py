"""Dump Kiwi summaries for batch-1 CGAP case ids."""
import csv
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BATCH = ROOT / "reports" / "easy_automation_batches.csv"
XLSX = ROOT / "docs" / "MoodleRegression_Testing_Kiwi new.xlsx"

ids = set()
with BATCH.open(encoding="utf-8") as f:
    for row in csv.DictReader(f):
        if row.get("batch") != "1":
            continue
        cid = row.get("case_id", "")
        if cid.startswith("CGAP-"):
            ids.add(cid)

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
for name in wb.sheetnames:
    if name.startswith("Regression") or name in ("Automation_Coverage", "Release_Regression_Plan"):
        continue
    ws = wb[name]
    for r in range(2, ws.max_row + 1):
        sno = str(ws.cell(r, 1).value or "").strip()
        if sno not in ids:
            continue
        title = ws.cell(r, 2).value
        summary = ws.cell(r, 3).value
        print(f"{name}\t{sno}\t{title}\t{summary}")
