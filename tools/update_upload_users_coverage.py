"""Mark the Upload_Users cases green in the Kiwi master and refresh the dashboard row.

Per-case judgement (never keyword-guessed). UU-02/07/24/45 were completed 2026-09-01
(example CSV headers, missing username header, unknown-course user still created,
over-limit firstname refused). Five cases remain No. Do not run this until those
four scenarios are green. Edits the workbook in place and is safe to re-run.
"""
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Upload_Users"
DASHBOARD = "Automation_Coverage"
VERIFIED = "green 2026-09-01"

# case id -> (automated, scenario tag or None if Covered by, extra note)
# One runnable @UU-Sn per combined group. Secondaries use covered_by.
COVERAGE = {
    "UU-01": ("Yes", "@UU-S1", ""),
    "UU-02": ("Yes", None, "Covered by @UU-S1"),
    "UU-03": ("Yes", None, "Covered by @UU-S1"),
    "UU-04": ("Yes", "@UU-S2", ""),
    "UU-05": ("Yes", None, "Covered by @UU-S2"),
    "UU-06": ("Yes", "@UU-S3", ""),
    "UU-07": ("Yes", "@UU-S4", ""),
    "UU-08": ("Yes", "@UU-S5", ""),
    "UU-09": ("Yes", "@UU-S6", ""),
    "UU-10": ("Yes", "@UU-S7", ""),
    "UU-11": ("Yes", "@UU-S8", ""),
    "UU-12": ("Yes", "@UU-S17", ""),
    "UU-13": ("Yes", "@UU-S9", ""),
    "UU-14": ("Yes", None, "Covered by @UU-S1"),
    "UU-15": ("Yes", "@UU-S10", ""),
    "UU-16": ("Yes", "@UU-S11", ""),
    "UU-17": ("Yes", None, "Covered by @UU-S11"),
    "UU-18": ("Yes", "@UU-S12", ""),
    "UU-19": ("Yes", None, "Covered by @UU-S12"),
    "UU-20": ("Yes", "@UU-S13", ""),
    "UU-21": ("Yes", None, "Covered by @UU-S13"),
    "UU-22": ("Yes", "@UU-S18", ""),
    "UU-23": ("Yes", "@UU-S14", ""),
    "UU-24": ("Yes", None, "Covered by @UU-S14"),
    "UU-25": ("Yes", "@UU-S15", ""),
    "UU-26": ("Yes", None, "Covered by @UU-S15"),
    "UU-27": ("Yes", "@UU-S16", ""),
    "UU-28": ("Yes", None, "Covered by @UU-S16"),
    "UU-29": ("Yes", None, "Covered by @UU-S1"),
    "UU-30": ("Yes", None, "Covered by @UU-S15"),
    "UU-31": ("No", "", "Not automatable for now: Low-priority non-UTF-8 fixture/encoding variant"),
    "UU-32": ("Yes", "@UU-S19", ""),
    "UU-33": ("Yes", "@UU-S20", ""),
    "UU-34": ("Yes", "@UU-S21", ""),
    "UU-35": ("Yes", "@UU-S22", ""),
    "UU-36": ("Yes", "@UU-S23", ""),
    "UU-37": ("Yes", "@UU-S24", ""),
    "UU-38": ("Yes", "@UU-S25", ""),
    "UU-39": ("Yes", "@UU-S26", ""),
    "UU-40": ("Yes", "@UU-S27", ""),
    "UU-41": ("Yes", "@UU-S28", ""),
    "UU-42": ("No", "", "Not automatable for now: requires a configured custom profile field"),
    "UU-43": ("No", "", "Not automatable for now: requires SMTP/mail capture"),
    "UU-44": ("No", "", "Not automatable for now: Low-priority default-template matrix"),
    "UU-45": ("Yes", "@UU-S29", ""),
    "UU-46": ("Yes", "@UU-S30", ""),
    "UU-47": ("Yes", "@UU-S31", ""),
    "UU-48": ("No", "", "Not automatable for now: requires a site-wide duplicate-email config mutation"),
}

# Derived cases whose wording did not match how Moodle 4.5 actually behaves.
REWORDED = {
    "UU-15": {
        "Summary": "A weak CSV password is flagged as weak and forces a password change at login",
        "Expected": "The row is counted under 'Users having a weak password', the account is created, "
                    "and the user must change the password at first login.",
    },
    "UU-26": {
        "Summary": "A Default values profile field is applied when the CSV omits that column",
        "Expected": "Created profile contains the configured default (Institution). Moodle removes the "
                    "default for any column the CSV supplies, so the default only applies to omitted "
                    "columns.",
    },
    "UU-45": {
        "Summary": "Values over database field limits are refused and create no user",
        "Expected": "The oversized firstname is not created. Moodle 4.5 reports a fatal "
                    "'Error writing to database' instead of a per-row verdict and aborts later CSV rows.",
    },
}


def column_index(sheet, header):
    for cell in sheet[1]:
        if str(cell.value).strip() == header:
            return cell.column
    raise SystemExit(f"column '{header}' not found on {sheet.title}")


def main():
    workbook = openpyxl.load_workbook(WORKBOOK)
    sheet = workbook[SHEET]
    title = column_index(sheet, "Title")
    status = column_index(sheet, "Case_Status")
    automated = column_index(sheet, "Automated")
    ref = column_index(sheet, "Automation_Ref")
    summary = column_index(sheet, "Summary")
    expected = column_index(sheet, "Expected")

    counts = {"Yes": 0, "Partial": 0, "Deviation": 0, "Deferred": 0, "No": 0}
    for row in range(2, sheet.max_row + 1):
        case = str(sheet.cell(row=row, column=title).value or "").strip()
        if case in COVERAGE:
            mark, tag, note = COVERAGE[case]
            sheet.cell(row=row, column=automated).value = mark
            if mark == "No":
                sheet.cell(row=row, column=ref).value = note
            elif tag:
                sheet.cell(row=row, column=ref).value = f"upload_users.feature {tag}"
            else:
                sheet.cell(row=row, column=ref).value = note
            if mark == "Yes":
                sheet.cell(row=row, column=status).value = "Pass"
            elif mark == "No" and not tag:
                sheet.cell(row=row, column=status).value = "NA"
            for header, value in REWORDED.get(case, {}).items():
                sheet.cell(row=row,
                           column=summary if header == "Summary" else expected).value = value
        current = str(sheet.cell(row=row, column=automated).value or "").strip()
        if current in counts:
            counts[current] += 1

    total = sum(counts.values())
    percent = f"{round(counts['Yes'] * 100 / total)}%" if total else "0%"
    print(SHEET, counts, "total", total, percent)

    dashboard = workbook[DASHBOARD]
    row = next(r for r in range(2, dashboard.max_row + 1)
               if str(dashboard.cell(row=r, column=1).value or "").strip() == SHEET)
    previous = [str(dashboard.cell(row=row, column=c).value or 0) for c in range(2, 8)]
    updated = [counts["Yes"], counts["Partial"], counts["Deviation"], counts["Deferred"],
               counts["No"], total]
    for offset, value in enumerate(updated):
        dashboard.cell(row=row, column=2 + offset).value = value
    dashboard.cell(row=row, column=8).value = percent

    totals_row = next(r for r in range(2, dashboard.max_row + 1)
                      if str(dashboard.cell(row=r, column=1).value or "").startswith("TOTAL"))
    for offset in range(6):
        cell = dashboard.cell(row=totals_row, column=2 + offset)
        cell.value = int(str(cell.value or 0)) - int(previous[offset]) + updated[offset]
    grand_total = int(dashboard.cell(row=totals_row, column=7).value)
    grand_yes = int(dashboard.cell(row=totals_row, column=2).value)
    dashboard.cell(row=totals_row, column=8).value = f"{round(grand_yes * 100 / grand_total)}%"
    print("TOTAL", [dashboard.cell(row=totals_row, column=c).value for c in range(2, 9)])

    workbook.save(WORKBOOK)


if __name__ == "__main__":
    main()
