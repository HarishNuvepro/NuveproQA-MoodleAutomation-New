"""
Stamp All Course Page + Moodle_Sandbox_Labs Automation_Ref to one runnable tag.

Primary row: feature @TAG
Other Yes rows in the same combined group: Covered by @TAG

Does not rewrite Automation_Coverage or case status.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "MoodleRegression_Testing_Kiwi new.xlsx"

ACP_SHEET = "All Course Page "
SB_SHEET = "Moodle_Sandbox_Labs"

# Excel row -> Automation_Ref  (Automated stays Yes)
ACP_REFS = {}


def acp_group(tag: str, rows: list[int]) -> None:
    ACP_REFS[rows[0]] = f"all_courses.feature {tag}"
    for r in rows[1:]:
        ACP_REFS[r] = f"Covered by {tag}"


acp_group("@ACP-S1", [4])
acp_group("@ACP-S2", [5, 6, 7, 8, 9, 54, 55, 56])
acp_group("@ACP-S3", [10, 11, 13, 29, 96])
acp_group("@ACP-S4", [14, 23])
acp_group("@ACP-S6", [17])
acp_group("@ACP-S7", [18])
acp_group("@ACP-S8", [21])
acp_group("@ACP-S9", [22])
acp_group("@ACP-S10", [24])
acp_group("@ACP-S11", [25, 26])
acp_group("@ACP-S12", [27, 28])
acp_group("@ACP-S13", [30, 31, 32, 33, 34, 35])
acp_group("@ACP-S14", list(range(37, 46)))
acp_group("@ACP-S15", [46, 47, 48, 49, 50, 51, 52, 53, 57, 58, 59])
acp_group("@ACP-S16", [60])
acp_group("@ACP-S17", [61, 62, 63, 64, 65, 66, 68, 69, 70, 71, 72, 83, 84])
acp_group("@ACP-S18", [73, 74])
acp_group("@ACP-S19", [75, 79, 80, 81, 82, 85])
acp_group("@ACP-S20", [86, 87, 88, 89, 90, 91, 92, 93, 101, 104])
acp_group("@ACP-S21", [95])
acp_group("@ACP-S23", [99, 100])
acp_group("@ACP-S24", [102])
acp_group("@ACP-S25", [103])
acp_group("@ACP-S26", [105])
acp_group("@ACP-S27", [106])
acp_group("@ACP-S29", [107])
acp_group("@CGAP-ACP-1", [108])
acp_group("@CGAP-ACP-2", [109])
acp_group("@CGAP-ACP-3", [110])
acp_group("@CGAP-ACP-4", [111])

SB_REFS = {}


def sb_group(feature: str, tag: str, rows: list[int]) -> None:
    SB_REFS[rows[0]] = f"{feature} {tag}"
    for r in rows[1:]:
        SB_REFS[r] = f"Covered by {tag}"


# create
sb_group("sandbox.feature", "@SB-create-1", [2])
sb_group("sandbox.feature", "@SB-create-9", [17])

# user (combined leftovers remapped to the surviving feature tag)
sb_group("sandbox_user_functionality.feature", "@SB-user-1", [7, 12, 13, 14, 25, 31, 44])
sb_group("sandbox_user_functionality.feature", "@SB-user-3", [15])
sb_group("sandbox_user_functionality.feature", "@SB-user-4", [41])
sb_group("sandbox_user_functionality.feature", "@SB-user-5", [26])
sb_group("sandbox_user_functionality.feature", "@SB-user-6", [9, 28])
sb_group("sandbox_user_functionality.feature", "@SB-user-7", [56, 57, 58, 59, 63])
sb_group("sandbox_user_functionality.feature", "@SB-user-8", [43])
sb_group("sandbox_user_functionality.feature", "@SB-user-11", [42])
sb_group("sandbox_user_functionality.feature", "@SB-user-12", [46])
sb_group("sandbox_user_functionality.feature", "@SB-user-13", [47])
sb_group("sandbox_user_functionality.feature", "@SB-user-14", [5, 8, 29, 51, 54])
sb_group("sandbox_user_functionality.feature", "@SB-user-18", [6])
sb_group("sandbox_user_functionality.feature", "@SB-user-19", [10, 39, 53])
sb_group("sandbox_user_functionality.feature", "@SB-manage-26", [40])
sb_group("sandbox_user_functionality.feature", "@SB-manage-27", [4, 27])
sb_group("sandbox_user_functionality.feature", "@SB-manage-29", [45])
sb_group("sandbox_user_functionality.feature", "@SB-manage-31", [16, 30])
sb_group("sandbox_user_functionality.feature", "@SB-manage-34", [32])

# manage
sb_group("sandbox.feature", "@SB-manage-1", [3, 21, 38])
sb_group("sandbox.feature", "@SB-manage-3", [33])
sb_group("sandbox.feature", "@SB-manage-8", [19, 22, 34, 48])
sb_group("sandbox.feature", "@SB-manage-9", [20, 35])
sb_group("sandbox.feature", "@SB-manage-12", [24])
sb_group("sandbox.feature", "@SB-manage-15", [23])
sb_group("sandbox.feature", "@SB-manage-20", [36])
sb_group("sandbox.feature", "@SB-manage-23", [11])


def apply(ws, mapping: dict[int, str]) -> int:
    changed = 0
    for row, ref in sorted(mapping.items()):
        current = str(ws.cell(row, 12).value or "")
        if current != ref:
            ws.cell(row, 12).value = ref
            changed += 1
        print(f"  {row:3} {ref}")
    return changed


def main():
    wb = load_workbook(XLSX)
    print("==== All Course Page ====")
    acp_n = apply(wb[ACP_SHEET], ACP_REFS)
    print("==== Moodle_Sandbox_Labs ====")
    sb_n = apply(wb[SB_SHEET], SB_REFS)
    tmp = XLSX.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    wb.close()
    tmp.replace(XLSX)

    check = load_workbook(XLSX, data_only=False)
    bad = []
    for sheet, mapping in ((ACP_SHEET, ACP_REFS), (SB_SHEET, SB_REFS)):
        ws = check[sheet]
        for row, ref in mapping.items():
            got = str(ws.cell(row, 12).value or "")
            if got != ref:
                bad.append((sheet, row, got, ref))
    check.close()
    print(f"Saved {XLSX} changed={acp_n + sb_n}")
    if bad:
        print("VERIFY FAILED")
        for item in bad[:20]:
            print(" ", item)
        raise SystemExit(1)
    print("VERIFY ok")


if __name__ == "__main__":
    main()
