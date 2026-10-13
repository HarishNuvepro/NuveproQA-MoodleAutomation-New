"""Mark Sequential_Activity_Completion Automated + Automation_Ref; repair CGAP row layout."""
from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Sequential_Activity_Completion"
COVERAGE = "Automation_Coverage"
PLAN = "Release_Regression_Plan"
FEATURE = "sequential_completion.feature"
STATES = ("yes", "partial", "deviation", "deferred", "no")

# Row 2 -> @SAC-1 … row 17 -> @SAC-16 (Kiwi row 11 = duplicate SAC-10 → @SAC-1)
YES_ROWS = list(range(2, 18))

CGAP_REFS = {
    f"CGAP-SAC-PP-{n}": f"{FEATURE} @CGAP-SAC-PP-{n}"
    for n in range(1, 13)
}

CGAP_PRIORITY = {
    "CGAP-SAC-PP-8": "High",
    "CGAP-SAC-PP-9": "High",
}
DEFAULT_CGAP_PRIORITY = "Medium"


def refresh_coverage_row(book) -> None:
    sheet = book[SHEET]
    counts = dict.fromkeys(STATES, 0)
    total = 0
    for row in range(2, sheet.max_row + 1):
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        total += 1
        state = str(sheet.cell(row, 11).value or "").strip().lower()
        if state in counts:
            counts[state] += 1

    coverage = book[COVERAGE]
    target = next(
        (row for row in range(2, coverage.max_row + 1)
         if str(coverage.cell(row, 1).value or "").strip() == SHEET),
        None,
    )
    if target is None:
        total_row = next(
            row for row in range(2, coverage.max_row + 1)
            if str(coverage.cell(row, 1).value or "").upper().startswith("TOTAL")
        )
        target = total_row
        coverage.insert_rows(target)
        template = target - 1
        for column in range(1, 9):
            coverage.cell(target, column)._style = copy(coverage.cell(template, column)._style)
        coverage.cell(target, 1).value = SHEET
        print(f"inserted {COVERAGE} row for {SHEET}")

    values = [counts[state] for state in STATES] + [total]
    for offset, value in enumerate(values, start=2):
        coverage.cell(target, offset).value = value
    coverage.cell(target, 8).value = (
        f"{round(100 * counts['yes'] / total)}%" if total else "0%"
    )
    print(f"{COVERAGE} {SHEET}: {values} (total={total})")


def repair_cgap_row(sheet, row: int, case_id: str) -> None:
    """Align CGAP rows with SAC column layout (Product/Category/Priority/Case_Status)."""
    priority = CGAP_PRIORITY.get(case_id, DEFAULT_CGAP_PRIORITY)
    sheet.cell(row, 4).value = "Moodle"
    sheet.cell(row, 5).value = "Sequential_Activity_Completion"
    sheet.cell(row, 6).value = priority
    sheet.cell(row, 7).value = "Pass"
    sheet.cell(row, 11).value = "Yes"
    sheet.cell(row, 12).value = CGAP_REFS[case_id]
    if not str(sheet.cell(row, 10).value or "").strip():
        sheet.cell(row, 10).value = "Practice Project"


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sheet = book[SHEET]
    sheet.cell(1, 11).value = "Automated"
    sheet.cell(1, 12).value = "Automation_Ref"

    for row in YES_ROWS:
        sac = row - 1
        sheet.cell(row, 11).value = "Yes"
        if sac == 10:
            sheet.cell(row, 12).value = (
                f"{FEATURE} @SAC-1 (duplicate SAC-10 removed — same steps as @SAC-1)"
            )
        else:
            sheet.cell(row, 12).value = f"{FEATURE} @SAC-{sac}"
        if not str(sheet.cell(row, 7).value or "").strip():
            sheet.cell(row, 7).value = "Pass"

    cgap_fixed = 0
    for row in range(2, sheet.max_row + 1):
        case_id = str(sheet.cell(row, 1).value or "").strip()
        if case_id in CGAP_REFS:
            repair_cgap_row(sheet, row, case_id)
            cgap_fixed += 1
    print(f"repaired {cgap_fixed} CGAP-SAC-PP rows")

    refresh_coverage_row(book)
    plan = book[PLAN]
    for row in range(10, plan.max_row + 1):
        if str(plan.cell(row, 2).value or "").strip() == SHEET:
            plan.cell(row, 11).value = "Yes"
            break

    book.save(WORKBOOK)
    print(f"saved {WORKBOOK}")


if __name__ == "__main__":
    main()
