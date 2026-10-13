"""Mark Randomization_Activity rows linked to randomactivity.feature (@RA-* / @CGAP-RAND-*)."""
from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
COVERAGE = "Automation_Coverage"
FEATURE = "randomactivity.feature"

# Row index on Randomization_Activity sheet (S.No in col A may differ for gap rows).
MAPPING: dict[int, tuple[str, str]] = {
    2: ("Yes", "@RA-create"),
    3: ("Yes", "@RA-name"),
    4: ("Yes", "@RA-name"),
    5: ("Partial", "@RA-name (same name validation scenario)"),
    70: ("Yes", "@CGAP-RAND-1 @RA-required"),
    71: ("Yes", "@CGAP-RAND-2 @RA-dupnames"),
    72: ("Yes", "@CGAP-RAND-3 @RA-form"),
}


def refresh_coverage_row(book, sheet_name: str) -> None:
    sheet = book[sheet_name]
    counts = {k: 0 for k in ("yes", "partial", "deviation", "deferred", "no")}
    total = 0
    for row in range(2, sheet.max_row + 1):
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        total += 1
        state = str(sheet.cell(row, 11).value or "").strip().lower()
        if state in counts:
            counts[state] += 1

    coverage = book[COVERAGE]
    target = next(
        (row for row in range(2, coverage.max_row + 1)
         if str(coverage.cell(row, 1).value or "").strip() == sheet_name),
        None,
    )
    if target is None:
        return
    values = [counts[k] for k in ("yes", "partial", "deviation", "deferred", "no")] + [total]
    for offset, value in enumerate(values, start=2):
        coverage.cell(target, offset).value = value
    auto = total - counts["no"]
    coverage.cell(target, 8).value = f"{round(100 * counts['yes'] / auto)}%" if auto else "0%"
    print(f"{COVERAGE} {sheet_name}: {values}")


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sheet = book["Randomization_Activity"]
    sheet.cell(1, 11).value = "Automated"
    sheet.cell(1, 12).value = "Automation_Ref"
    for row, (auto, ref) in MAPPING.items():
        sheet.cell(row, 11).value = auto
        sheet.cell(row, 12).value = f"{FEATURE} {ref}"
        if str(sheet.cell(row, 7).value or "").strip().lower() in ("confirmed", "pass", "proposed"):
            sheet.cell(row, 7).value = "Pass"
    refresh_coverage_row(book, "Randomization_Activity")
    book.save(WORKBOOK)
    print(f"saved {WORKBOOK} ({len(MAPPING)} rows updated)")


if __name__ == "__main__":
    main()
