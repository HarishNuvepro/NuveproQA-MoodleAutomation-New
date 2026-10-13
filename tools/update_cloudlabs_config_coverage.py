"""Mark Difficulty_Level_Configuration + Content_Integration_Git_URL in Kiwi."""
from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
COVERAGE = "Automation_Coverage"
PLAN = "Release_Regression_Plan"
STATES = ("yes", "partial", "deviation", "deferred", "no")
FEATURE_GIT = "cloudlabs_config.feature"
FEATURE_DIFFICULTY = "difficulty_level_configuration.feature"

SHEETS = {
    "Difficulty_Level_Configuration": {
        2: ("Yes", "@DLC-1"),
        3: ("Yes", "@DLC-2"),
        4: ("Deviation", "@DLC-3 (live levels differ from Practitioner/Proficient doc)"),
        5: ("Yes", "@DLC-4"),
        6: ("Yes", "@DLC-5"),
        7: ("Yes", "@DLC-6"),
        8: ("Yes", "@DLC-10"),
        9: ("Yes", "@DLC-8"),
        10: ("Yes", "@DLC-9"),
        11: ("Yes", "@DLC-11"),
        12: ("Yes", "@DLC-12 (duplicate preserves difficulty; full cross-Moodle mbz out of scope)"),
        13: ("Yes", "@DLC-13 (admin restore default four levels → new activity form)"),
        14: ("Yes", "@DLC-14 (remove in-use level — edit form loads; dropdown remaps to remaining level)"),
        # Row 15 (report filter gap) removed from sheet — not needed.
        15: ("Yes", "@DLC-5/@DLC-6 (admin→activity flow)"),
        16: ("Yes", "@DLC-4 empty config fallback"),
        17: ("Deviation", "@DLC-3"),
    },
    "Content_Integration_Git_URL": {
        2: ("Yes", "@CIG-1"),
        3: ("Yes", "@CIG-2"),
        4: ("Yes", "@CIG-3"),
        5: ("No", "Valid git fetch — needs network + stable URL"),
        6: ("No", "Invalid git fetch — needs network"),
        7: ("Deferred", "Switch manual/git after save — not scripted"),
        8: ("Deviation", "@CIG-7 @deviation — playground still shows description source"),
        9: ("Yes", "@CIG-10 guided content source"),
        10: ("Partial", "@CIG-10 manual+git toggle only"),
        11: ("Yes", "@CIG-11 assessment content source"),
        12: ("Partial", "@CIG-11 manual+git toggle only"),
        13: ("No", "Rendered description content — needs git fetch"),
        14: ("No", "Duplicate activity content — not scripted"),
        15: ("No", "Live git sync — not scripted"),
        16: ("No", ".md/.txt git formats — needs git"),
        17: ("No", "Mixed manual+git save — not scripted"),
        18: ("No", "Images in git markdown — needs git"),
        19: ("No", "Parallel manual/git view — not scripted"),
        20: ("No", "PDF export content — not scripted"),
        21: ("No", "Randomized PS git — needs assessment fixture"),
        22: ("No", "Screenshot in git — needs git"),
        23: ("No", "Backup/restore content — not scripted"),
        24: ("No", "Offline git cache — not scripted"),
        25: ("No", "Private git auth — not scripted"),
        26: ("Deferred", "Invalid git URL error — no safe fetch hook"),
        27: ("Partial", "@CIG-2/@CIG-3 cross lab types"),
    },
}


def refresh_coverage_row(book, sheet_name: str) -> None:
    sheet = book[sheet_name]
    counts = dict.fromkeys(STATES, 0)
    total = 0
    for row in range(2, sheet.max_row + 1):
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        total += 1
        state = str(sheet.cell(row, 11).value or "").strip().lower()
        if state in counts:
            counts[state] += 1

    coverage = book[COVERAGE]
    target = next(
        (row for row in range(2, coverage.max_row + 1)
         if str(coverage.cell(row, 1).value or "").strip() == sheet_name),
        None,
    )
    if target is None:
        total_row = next(
            row for row in range(2, coverage.max_row + 1)
            if str(coverage.cell(row, 1).value or "").upper().startswith("TOTAL")
        )
        target = total_row
        coverage.insert_rows(target)
        template = target - 1
        for column in range(1, 9):
            coverage.cell(target, column)._style = copy(coverage.cell(template, column)._style)
        coverage.cell(target, 1).value = sheet_name
        print(f"inserted {COVERAGE} row for {sheet_name}")

    values = [counts[state] for state in STATES] + [total]
    for offset, value in enumerate(values, start=2):
        coverage.cell(target, offset).value = value
    coverage.cell(target, 8).value = (
        f"{round(100 * counts['yes'] / total)}%" if total else "0%"
    )
    print(f"{COVERAGE} {sheet_name}: {values}")


def apply_sheet(book, sheet_name: str, mapping: dict, feature_file: str) -> None:
    sheet = book[sheet_name]
    sheet.cell(1, 11).value = "Automated"
    sheet.cell(1, 12).value = "Automation_Ref"
    for row, (auto, ref) in mapping.items():
        sheet.cell(row, 11).value = auto
        sheet.cell(row, 12).value = f"{feature_file} {ref}"
        if str(sheet.cell(row, 7).value or "").strip().lower() in ("confirmed", "pass", "proposed"):
            sheet.cell(row, 7).value = "Pass"
    refresh_coverage_row(book, sheet_name)


def mark_plan(book, sheet_name: str) -> None:
    plan = book[PLAN]
    for row in range(10, plan.max_row + 1):
        if str(plan.cell(row, 2).value or "").strip() == sheet_name:
            plan.cell(row, 11).value = "Yes"
            print(f"{PLAN} Automated=Yes for {sheet_name}")
            return


def delete_difficulty_report_filter_row(book) -> None:
    """Remove obsolete row 15 (report filter gap) once; idempotent if already deleted."""
    sheet = book["Difficulty_Level_Configuration"]
    for row in range(2, sheet.max_row + 1):
        sno = str(sheet.cell(row, 1).value or "").strip()
        ref = str(sheet.cell(row, 12).value or "").lower()
        if "report filter" in ref or (
            row == 15 and "gap check" in sno.lower() and "report" in str(sheet.cell(row, 3).value or "").lower()
        ):
            sheet.delete_rows(row, 1)
            print(f"deleted Difficulty_Level_Configuration row {row} (report filter — not needed)")
            return
    # Legacy: fixed row 15 before CGAP shift
    if sheet.max_row >= 15:
        title = str(sheet.cell(15, 3).value or "").lower()
        if "report" in title and "filter" in title:
            sheet.delete_rows(15, 1)
            print("deleted Difficulty_Level_Configuration row 15 (report filter)")


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    delete_difficulty_report_filter_row(book)
    feature_for = {
        "Difficulty_Level_Configuration": FEATURE_DIFFICULTY,
        "Content_Integration_Git_URL": FEATURE_GIT,
    }
    for sheet_name, mapping in SHEETS.items():
        apply_sheet(book, sheet_name, mapping, feature_for[sheet_name])
        mark_plan(book, sheet_name)
    book.save(WORKBOOK)
    print(f"saved {WORKBOOK}")


if __name__ == "__main__":
    main()
