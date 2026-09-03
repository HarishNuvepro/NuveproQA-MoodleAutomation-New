"""Move leftover Lab Templates Partials to No. Leaves Deviations."""
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
ROWS = {
    47: "No — browser errors covered by @LT-46; server-log verification needs log access",
    70: "No — Sandbox Create/Start covered by @LT-69; Guided/Assessment/Unpark need extra fixtures",
    78: "No — Create-failure recovery covered by @LT-77; Start/Unpark failure injection unavailable",
}


def main() -> None:
    book = load_workbook(XLSX)
    sheet = book["Moodle_Lab_Templates"]
    for row, reason in ROWS.items():
        sheet.cell(row, 11).value = "No"
        sheet.cell(row, 12).value = reason
        sheet.cell(row, 7).value = "NA"
        print(f"  {row} No")
    book.save(XLSX)
    print("saved", XLSX)


if __name__ == "__main__":
    main()
