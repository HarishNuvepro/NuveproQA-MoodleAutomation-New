"""
Refresh Moodle_Sandbox_Labs Automated + Automation_Ref for converted sandbox cases.
Updates the workbook in place. Touch only that sheet + Automation_Coverage totals.
"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "MoodleRegression_Testing_Kiwi new.xlsx"
SHEET = "Moodle_Sandbox_Labs"

# Match against Summary (case-insensitive). First match wins.
PATCHES = [
    ("loading in i frame",
     "Yes",
     "sandbox_user_functionality.feature @SB-user-14"),
    ("alloted,consumed and remaining timing",
     "Yes",
     "sandbox_user_functionality.feature @SB-user-18"),
    ("speed test button is working",
     "Yes",
     "Covered by @SB-user-14"),
    ("catalog ui input is enabled",
     "Yes",
     "sandbox_user_functionality.feature @SB-user-19"),
    ("speed test runs and displays a result",
     "Yes",
     "Covered by @SB-user-14"),
    ("lab details' button",
     "No",
     "Not automatable — Lab Details full field set needs a running lab plus Manage Labs details"),
    ("invalid actions are prevented",
     "Yes",
     "Covered by @SB-user-19"),
    ("no access link",
     "No",
     "Not automatable — no-access-link / incorrect-config needs a specially broken lab fixture"),
    ("launched in another tab",
     "Yes",
     "Covered by @SB-user-14"),
    ("nears its duration expiry",
     "No",
     "Not automatable — near-expiry warning needs notification+duration soak"),
    ("unlimited-duration template does not expire",
     "Yes",
     "Covered by @SB-user-19"),
    ("automatically expires and is stopped",
     "No",
     "Not automatable — time-based auto-expire soak"),
    ("duration expired' message",
     "No",
     "Not automatable — expired-duration fixture / message"),
    ("description is displaying in sandbox landing",
     "Yes",
     "sandbox_user_functionality.feature @SB-user-7"),
    ("lab type, lab status",
     "Yes",
     "Covered by @SB-user-7"),
    ("duration is displaying in sandbox landing",
     "Yes",
     "Covered by @SB-user-7"),
    ("sandbox id is displaying",
     "Yes",
     "Covered by @SB-user-7"),
    ("no lab exists yet",
     "Yes",
     "Covered by @SB-user-7"),
    ("loading full screen",
     "Yes",
     "Covered by @SB-user-14"),
]

FULLSCREEN_SUMMARY = (
    "Verify Vm lab is loading full screen."
)


def recount(ws):
    c = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        auto = str(row[10]).strip() if len(row) > 10 and row[10] is not None else ""
        if auto.startswith("No"):
            key = "No"
        elif auto in ("Yes", "Partial", "Deviation", "Deferred"):
            key = auto
        else:
            key = "No"
        c[key] += 1
    return c


def update_dashboard(wb, counts):
    ws = wb["Automation_Coverage"]
    total = sum(counts.values())
    yes = counts.get("Yes", 0)
    pct = f"{round(100 * yes / total)}%" if total else "0%"
    sandbox_row = None
    total_row = None
    for row in ws.iter_rows(min_row=2, max_row=30):
        name = row[0].value
        if name == SHEET:
            sandbox_row = row
        elif name and str(name).startswith("TOTAL"):
            total_row = row
    if sandbox_row is None:
        raise RuntimeError(f"{SHEET} missing on Automation_Coverage")
    sandbox_row[1].value = yes
    sandbox_row[2].value = counts.get("Partial", 0)
    sandbox_row[3].value = counts.get("Deviation", 0)
    sandbox_row[4].value = counts.get("Deferred", 0)
    sandbox_row[5].value = counts.get("No", 0)
    sandbox_row[6].value = total
    sandbox_row[7].value = pct

    if total_row is not None:
        acc = Counter()
        for row in ws.iter_rows(min_row=2, max_row=17):
            if row[0].value is None or str(row[0].value).startswith("TOTAL"):
                continue
            acc["Yes"] += int(row[1].value or 0)
            acc["Partial"] += int(row[2].value or 0)
            acc["Deviation"] += int(row[3].value or 0)
            acc["Deferred"] += int(row[4].value or 0)
            acc["No"] += int(row[5].value or 0)
        t = sum(acc.values())
        total_row[1].value = acc["Yes"]
        total_row[2].value = acc["Partial"]
        total_row[3].value = acc["Deviation"]
        total_row[4].value = acc["Deferred"]
        total_row[5].value = acc["No"]
        total_row[6].value = t
        total_row[7].value = f"{round(100 * acc['Yes'] / t)}%" if t else "0%"


def summary_of(ws, r):
    v = ws.cell(r, 2).value
    return "" if v is None else str(v)


def delete_duplicate_duration(ws):
    dropped = []
    for r in range(ws.max_row, 1, -1):
        s = summary_of(ws, r).lower()
        if "allotted / consumed / remaining duration values are shown correctly" in s:
            ws.delete_rows(r, 1)
            dropped.append(r)
    return dropped


def append_fullscreen(ws):
    for r in range(2, ws.max_row + 1):
        if "loading full screen" in summary_of(ws, r).lower():
            return False
    src = ws.max_row
    ws.insert_rows(src + 1)
    dest = src + 1
    for c in range(1, 13):
        ws.cell(dest, c).value = ws.cell(src, c).value
    ws.cell(dest, 1).value = "Sandbox Gap (automation-derived)"
    ws.cell(dest, 2).value = FULLSCREEN_SUMMARY
    ws.cell(dest, 3).value = "VM sandbox activity with running lab; Switch to full screen on console."
    ws.cell(dest, 6).value = "Medium"
    ws.cell(dest, 7).value = "PROPOSED"
    ws.cell(dest, 8).value = "Click Switch to full screen on the VM lab console and verify the lab loads full screen."
    ws.cell(dest, 9).value = "Lab loads in full screen."
    ws.cell(dest, 11).value = "Yes"
    ws.cell(dest, 12).value = (
        "Covered by @SB-user-14"
    )
    return True


def apply_patches(ws):
    applied = []
    for r in range(2, ws.max_row + 1):
        s = summary_of(ws, r).lower()
        if not s:
            continue
        for needle, status, ref in PATCHES:
            if needle in s:
                ws.cell(r, 11).value = status
                ws.cell(r, 12).value = ref
                applied.append((r, status, needle))
                break
    return applied


def main():
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    dropped = delete_duplicate_duration(ws)
    print("Deleted duplicate duration rows:", dropped)
    added = append_fullscreen(ws)
    print("Appended fullscreen row:", added)
    applied = apply_patches(ws)
    for r, status, needle in applied:
        print(f"{status:8} row {r}: {needle}")
    counts = recount(ws)
    update_dashboard(wb, counts)
    tmp = XLSX.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    wb.close()
    tmp.replace(XLSX)
    check = load_workbook(XLSX, data_only=False)
    cws = check[SHEET]
    c = recount(cws)
    print("Saved", XLSX)
    print("Counts", dict(c), "total", sum(c.values()),
          f"Yes%={round(100 * c.get('Yes', 0) / sum(c.values()))}%" if sum(c.values()) else "")
    for r in range(2, cws.max_row + 1):
        print(f"  {r} {cws.cell(r, 11).value}: {str(cws.cell(r, 2).value)[:80]}")
    check.close()


if __name__ == "__main__":
    main()
