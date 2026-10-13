"""Sync Pre_Delivery_Lab_Checklist Automated / Automation_Ref from predelivery*.feature."""
from __future__ import annotations

from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Pre_Delivery_Lab_Checklist"
COVERAGE = "Automation_Coverage"

YES: dict[int, str] = {
    2: "predelivery.feature @PDL-role",
    4: "predelivery.feature @PDL-progress",
    5: "predelivery.feature @PDL-validation",
    13: "predelivery.feature @CGAP-PDL-2",
    15: "predelivery.feature @PDL-sections",
    16: "predelivery.feature @PDL-catalog",
    17: "predelivery.feature @PDL-guided",
    25: "predelivery.feature @PDL-labtype",
    26: "predelivery.feature @PDL-student-tab",
    28: "predelivery.feature @PDL-progress",
    74: "predelivery.feature @CGAP-PDL-1",
    75: "predelivery.feature @CGAP-PDL-2",
    76: "predelivery.feature @CGAP-PDL-3",
    6: "predelivery_mutate.feature @PDL-save-first",
    9: "predelivery_mutate.feature @PDL-save-first",
    10: "predelivery_mutate.feature @PDL-audit-modify",
    7: "predelivery_mutate.feature @PDL-visibility",
    8: "predelivery_mutate.feature @PDL-visibility",
    11: "predelivery_mutate.feature @PDL-screenshot",
    12: "predelivery_mutate.feature @PDL-support-clear",
    77: "predelivery_mutate.feature @CGAP-PDL-4",
    19: "predelivery_dashboard.feature @PDL-dash-ui",
    20: "predelivery_dashboard.feature @PDL-dash-search",
    21: "predelivery_dashboard.feature @PDL-dash-filters",
    22: "predelivery_dashboard.feature @PDL-dash-filters",
    23: "predelivery_dashboard.feature @PDL-dash-filters",
    24: "predelivery_dashboard.feature @PDL-dash-open",
    30: "predelivery_dashboard.feature @PDL-dash-ui",
    31: "predelivery_dashboard.feature @PDL-dash-ui",
    32: "predelivery_dashboard.feature @PDL-dash-ui",
    33: "predelivery_dashboard.feature @PDL-dash-ui",
    35: "predelivery_dashboard.feature @PDL-dash-ui",
    36: "predelivery_dashboard.feature @PDL-dash-search",
    37: "predelivery_dashboard.feature @PDL-dash-search",
    38: "predelivery_dashboard.feature @PDL-dash-empty",
    39: "predelivery_dashboard.feature @PDL-dash-search",
    40: "predelivery_dashboard.feature @PDL-dash-search",
    41: "predelivery_dashboard.feature @PDL-dash-filters",
    42: "predelivery_dashboard.feature @PDL-dash-filters",
    44: "predelivery_dashboard.feature @PDL-dash-filters",
    45: "predelivery_dashboard.feature @PDL-dash-filters",
    46: "predelivery_dashboard.feature @PDL-dash-filters",
    47: "predelivery_dashboard.feature @PDL-dash-filters",
    48: "predelivery_dashboard.feature @PDL-dash-empty",
    49: "predelivery_dashboard.feature @PDL-dash-filters",
    50: "predelivery_dashboard.feature @PDL-dash-filters",
    51: "predelivery_dashboard.feature @PDL-dash-filters",
    52: "predelivery_dashboard.feature @PDL-dash-filters",
    55: "predelivery_dashboard.feature @PDL-dash-filters",
    56: "predelivery_dashboard.feature @PDL-dash-filters",
    57: "predelivery_dashboard.feature @PDL-dash-empty",
    58: "predelivery_dashboard.feature @PDL-dash-ui",
    59: "predelivery_dashboard.feature @PDL-dash-ui",
    60: "predelivery_dashboard.feature @PDL-dash-ui",
    61: "predelivery_dashboard.feature @PDL-dash-open",
    62: "predelivery_dashboard.feature @PDL-dash-open",
    63: "predelivery_dashboard.feature @PDL-dash-ui",
    64: "predelivery_dashboard.feature @PDL-dash-apply",
    65: "predelivery_dashboard.feature @PDL-dash-apply",
    67: "predelivery_dashboard.feature @PDL-67",
    69: "predelivery_dashboard.feature @PDL-69",
    70: "predelivery_dashboard.feature @PDL-69",
}

PARTIAL: dict[int, str] = {
    14: "Partial — force-enable below 100%; no stable selector documented",
    18: "Partial — visibility + feature-flag integration beyond single save",
    34: "Partial — card counts after live save; covered indirectly by @PDL-dash-ui read-only",
    43: "Partial — all courses in dropdown; needs env with many courses",
    53: "Partial — unassigned support engineer '-' filter; needs seeded row",
    54: "Partial — support engineer dropdown completeness; needs seeded data",
    71: "Partial — zero-lab dashboard; needs empty env",
    72: "Partial — force-enable gap; same as row 14",
    73: "Partial — dashboard count accuracy after save; combine mutate + dashboard runs",
}

NO: dict[int, str] = {
    3: "No — Plan/Team/Resource auto-populate from CloudLabs backend",
    27: "No — concurrent multi-user checklist edits; manual",
    29: "No — infinite scroll batch dedupe; heavy seed",
    66: "No — responsive layout; manual visual",
    68: "No — 2476-row performance; manual/perf env",
}


def refresh_coverage(book) -> None:
    sh = book[SHEET]
    counts = dict.fromkeys(("yes", "partial", "deviation", "deferred", "no"), 0)
    total = 0
    for row in range(2, sh.max_row + 1):
        if not str(sh.cell(row, 1).value or "").strip():
            continue
        total += 1
        st = str(sh.cell(row, 11).value or "").strip().lower()
        if st.startswith("no"):
            counts["no"] += 1
        elif st in counts:
            counts[st] += 1
        else:
            counts["no"] += 1
    cov = book[COVERAGE]
    for r in range(2, cov.max_row + 1):
        if str(cov.cell(r, 1).value or "").strip() == SHEET:
            cov.cell(r, 2).value = counts["yes"]
            cov.cell(r, 3).value = counts["partial"]
            cov.cell(r, 4).value = counts["deviation"]
            cov.cell(r, 5).value = counts["deferred"]
            cov.cell(r, 6).value = counts["no"]
            cov.cell(r, 7).value = total
            auto = total - counts["no"]
            cov.cell(r, 8).value = f"{round(100 * counts['yes'] / auto)}%" if auto else "0%"
            break


def apply() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sh = book[SHEET]

    def set_row(row: int, auto: str, ref: str, status: str) -> None:
        if row > sh.max_row or not str(sh.cell(row, 1).value or "").strip():
            return
        sh.cell(row, 11).value = auto
        sh.cell(row, 12).value = ref
        sh.cell(row, 7).value = status

    for row, ref in YES.items():
        set_row(row, "Yes", ref, "Pass")
    for row, ref in PARTIAL.items():
        set_row(row, "Partial", ref, "Blocked")
    for row, ref in NO.items():
        set_row(row, "No", ref, "NA")

    refresh_coverage(book)
    book.save(WORKBOOK)
    print(f"updated {WORKBOOK} — {SHEET}")


if __name__ == "__main__":
    apply()
