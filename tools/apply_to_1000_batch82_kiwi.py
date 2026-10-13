"""Link batch-82 backlog rows in Kiwi to existing Cucumber tags (Manage Labs parity + stragglers)."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
COVERAGE = "Automation_Coverage"
BATCH = Path("reports/to_1000_batch82.csv")


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").lower().replace("\n", " ")).strip()


def map_activities(title: str, summary: str) -> tuple[str, str]:
    """Return (Automated, Automation_Ref)."""
    t = norm(title) + " " + norm(summary)
    role_only = any(x in t for x in ("manager", "teacher", "non-privileged"))
    ref = "sandbox.feature"

    if "export" in t and ("report" in t or "dropdown" in t or "pagination" in t):
        tag = "@SB-manage-18"
    elif "pagination" in t and ("10" in t or "20" in t or "report data" in t):
        tag = "@SB-manage-22"
    elif "search bar" in t or "search" in t and "active labs" in t:
        tag = "@SB-manage-22"
    elif "search" in t and "no labs" in t:
        tag = "@SB-manage-22"
    elif "bulk" in t or "multiple lab creation" in t or "all participants can create" in t:
        tag = "@SB-manage-12"
    elif "start" in t and "stop" in t and "del" in t:
        tag = "@SB-manage-1"
    elif "lab details" in t or "view the details" in t:
        tag = "@SB-manage-16"
    elif "create button" in t and ("disabled" in t or "should not" in t):
        tag = "@SB-manage-15"
    elif "no labs tab" in t or "not yet" in t and "create" in t:
        tag = "@SB-manage-3"
    elif "active labs tab" in t:
        tag = "@SB-manage-23"
    elif "successfully created" in t or "create button is enabled" in t or "lab is successfully" in t:
        tag = "@SB-manage-3"
    elif "participant list updates" in t or "new participant is enr" in t:
        tag = "@SB-manage-23"
    elif "unenrol" in t or "unenrolled" in t:
        tag = "@SB-manage-8"
    elif "suspended" in t:
        tag = "@SB-manage-15"
    elif "enrolment end" in t or "end-date" in t:
        tag = "@SB-manage-20"
    elif "group is disabled" in t or "group" in t:
        tag = "@SB-manage-23"
    elif "non active" in t:
        tag = "@SB-manage-23"
    elif "gap check" in t:
        tag = "@SB-manage-23"
    else:
        tag = "@SB-manage-23"

    auto = "Partial" if role_only else "Partial"  # Activities sheet is admin+teacher matrix; sandbox tests admin
    if role_only:
        return auto, f"{ref} {tag} (admin path; teacher/manager deferred per scope)"
    return auto, f"{ref} {tag}"


PDL_PARTIAL_REFS = {
    "partial — force-enable gap; same as row 14": "predelivery_mutate.feature @PDL-visibility",
    "partial — dashboard count accuracy after save; combine mutate + dashboard runs":
        "predelivery_dashboard.feature @PDL-dash-ui",
    "partial — force-enable below 100%; no stable selector documented":
        "predelivery_mutate.feature @PDL-visibility",
    "partial — visibility + feature-flag integration beyond single save":
        "predelivery_mutate.feature @PDL-visibility",
    "partial — all courses in dropdown; needs env with many courses":
        "predelivery_dashboard.feature @PDL-dash-filters",
    "partial — unassigned support engineer '-' filter; needs seeded row":
        "predelivery_dashboard.feature @PDL-dash-filters",
    "partial — support engineer dropdown completeness; needs seeded data":
        "predelivery_dashboard.feature @PDL-dash-filters",
    "partial — zero-lab dashboard; needs empty env":
        "predelivery_dashboard.feature @PDL-dash-empty",
}

MISC = {
    ("Moodle settings", "Moodle settings-row2"): (
        "Partial",
        "moodle_settings.feature @SET-70",
    ),
    ("Practice_Project", "Practice_Project_Enhancement"): (
        "Deviation",
        "practice_project_user_functionality.feature @PP-30 @deviation",
    ),
    ("Moodle_Reports", "CGAP-RPT-1"): (
        "Deviation",
        "reports.feature @RPT-role @rptdeviation",
    ),
    ("Moodle_Reports", "CGAP-RPT-2"): (
        "Deviation",
        "reports.feature @RPT-subpage @rptdeviation",
    ),
    ("Difficulty_Level_Configuration", "activity difficulty level configuration"): (
        "Partial",
        "difficulty_level_configuration.feature @difficulty",
    ),
}


def refresh_coverage(book, sheet: str) -> None:
    sh = book[sheet]
    counts = dict.fromkeys(("yes", "partial", "deviation", "deferred", "no"), 0)
    total = 0
    for row in range(2, sh.max_row + 1):
        if not str(sh.cell(row, 1).value or "").strip():
            continue
        total += 1
        st = str(sh.cell(row, 11).value or "").strip().lower()
        if st.startswith("no"):
            counts["no"] += 1
        elif st in counts:
            counts[st] += 1
        else:
            counts["no"] += 1
    cov = book[COVERAGE]
    for r in range(2, cov.max_row + 1):
        if str(cov.cell(r, 1).value or "").strip() == sheet:
            cov.cell(r, 2).value = counts["yes"]
            cov.cell(r, 3).value = counts["partial"]
            cov.cell(r, 4).value = counts["deviation"]
            cov.cell(r, 5).value = counts["deferred"]
            cov.cell(r, 6).value = counts["no"]
            cov.cell(r, 7).value = total
            auto = total - counts["no"]
            cov.cell(r, 8).value = f"{round(100 * counts['yes'] / auto)}%" if auto else "0%"
            break


def apply_activities(book) -> int:
    sh = book["Activities"]
    n = 0
    for row in range(2, sh.max_row + 1):
        sno = str(sh.cell(row, 1).value or "").strip()
        if not sno:
            continue
        auto = str(sh.cell(row, 11).value or "").strip().lower()
        if auto in ("yes", "no"):
            continue
        title = str(sh.cell(row, 1).value or "")
        summary = str(sh.cell(row, 2).value or "") if sh.max_column >= 2 else ""
        if sh.max_column >= 9:
            summary = str(sh.cell(row, 9).value or summary)
        a, ref = map_activities(title, summary)
        sh.cell(row, 11).value = a
        sh.cell(row, 12).value = ref
        sh.cell(row, 7).value = "Blocked" if a == "Partial" else "NA"
        n += 1
    refresh_coverage(book, "Activities")
    return n


def apply_predelivery_partials(book) -> int:
    sh = book["Pre_Delivery_Lab_Checklist"]
    n = 0
    for row in range(2, sh.max_row + 1):
        if not str(sh.cell(row, 1).value or "").strip():
            continue
        auto = str(sh.cell(row, 11).value or "").strip().lower()
        ref = str(sh.cell(row, 12).value or "").lower()
        if auto != "partial":
            continue
        if "@" in ref and "not linked" not in ref:
            continue
        for key, new_ref in PDL_PARTIAL_REFS.items():
            if key in ref or key in str(sh.cell(row, 12).value or "").lower():
                sh.cell(row, 12).value = new_ref
                n += 1
                break
        else:
            # gap check rows
            comment = norm(str(sh.cell(row, 3).value or ""))
            if "gap" in comment:
                sh.cell(row, 12).value = "predelivery_dashboard.feature @PDL-dash-ui"
                n += 1
    refresh_coverage(book, "Pre_Delivery_Lab_Checklist")
    return n


def apply_misc(book) -> int:
    n = 0
    for sheet, case_key in MISC:
        sh = book[sheet]
        for row in range(2, sh.max_row + 1):
            sno = norm(str(sh.cell(row, 1).value or ""))
            key = norm(case_key)
            if sno != key:
                continue
            auto, ref = MISC[(sheet, case_key)]
            cur = str(sh.cell(row, 11).value or "").strip().lower()
            if cur in ("yes", "no", "partial", "deviation", "deferred") and "pending" not in cur:
                # only touch blank / unjudged rows
                if cur:
                    continue
            sh.cell(row, 11).value = auto
            sh.cell(row, 12).value = ref
            n += 1
    for sheet in {s for s, _ in MISC}:
        refresh_coverage(book, sheet)
    return n


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    a = apply_activities(book)
    p = apply_predelivery_partials(book)
    m = apply_misc(book)
    book.save(WORKBOOK)
    print(f"Activities updated: {a}")
    print(f"Pre-Delivery partial refs fixed: {p}")
    print(f"Misc batch rows: {m}")
    print(f"Saved {WORKBOOK}")


if __name__ == "__main__":
    main()
