"""Record exact verified Lab Templates coverage in the Kiwi master."""
from pathlib import Path

import openpyxl

WORKBOOK = Path("docs/MoodleRegression_Testing_Kiwi new.xlsx")
SHEET = "Moodle_Lab_Templates"

# Workbook row -> exact scenario reference verified green.
LABTEMPLATE_TAGS = (
    1, 2, 3, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24,
    25, 26, 27, 28, 29, 30, 31, 32, 33, 34, 35, 36, 38, 40, 41,
    42, 43, 44, 45, 47, 48, 49, 50, 51, 52, 53, 54, 55, 56, 57,
    58, 62, 76, 80,
)
VERIFIED = {
    tag + 1: f"labtemplates.feature @LT-{tag}"
    for tag in LABTEMPLATE_TAGS
}
VERIFIED.update({
    5: "sandbox_user_functionality.feature @LT-4",
    6: "sandbox_user_functionality.feature @LT-5",
    7: "assessment_user_functionality.feature @LT-6",
    8: "sandbox_user_functionality.feature @LT-7; sandbox.feature @LT-7",
    9: "sandbox_user_functionality.feature @LT-8",
    11: "sandbox_user_functionality.feature @LT-10",
    60: "sandbox_user_functionality.feature @LT-59",
    61: "sandbox_user_functionality.feature @LT-60",
    62: "sandbox_user_functionality.feature @LT-61",
    67: "sandbox_user_functionality.feature @LT-66",
    68: "sandbox.feature @LT-67",
    74: "sandbox.feature @LT-73",
    71: "sandbox.feature @LT-70",
    84: "labtemplates.feature @CGAP-LT-3",
    85: "labtemplates.feature @CGAP-LT-4",
})

DEVIATIONS = {}

PARTIAL = {}

DEFERRED = {}

BLOCKED = {
    47: "No — browser errors covered by @LT-46; server-log verification needs log access",
    70: "No — Sandbox Create/Start covered by @LT-69; Guided/Assessment/Unpark need extra fixtures",
    78: "No — Create-failure recovery covered by @LT-77; Start/Unpark failure injection unavailable",
    64: "Real Create timing outcome is unspecified and needs a dedicated slow-plan fixture",
    65: "Real Start timing outcome is unspecified and needs a dedicated slow-plan fixture",
    66: "No Park/Unpark API and no parked VM fixture",
    69: "CloudLabs client and Moodle UI expose Stop, but no Park action",
    72: "Proctored progress scenario is scripted but needs a working proctoring launch fixture",
    73: "Requires Orgo suspend/resume access; no API or UI automation endpoint is available",
    75: "Requires SSO identity-provider credentials and login flow",
    76: "Requires an LTI consumer and signed launch flow",
    79: "Requires a lab-create webhook inspection endpoint or server-log access",
    80: "Low priority and requires mutating/restoring the site-wide theme during transient progress",
    10: "No — @LT-9 @skip (sandbox_user_functionality): duration notification not verifiable; "
        "remove scenario + LabTemplateSteps/SB-user notification steps when dropped",
    83: "No — unbounded interval 999999999 is accepted; no max validation",
    38: "No — Randomization problem-statement Plan/Team/Catalog not automatable for now (manual @LT-37)",
    40: "No — Imported/restored course Plan/Team/Catalog not automatable for now (manual @LT-39)",
}

# Excel row 82 removed 2026-09-21 — bulk export/import not applicable to this application.
# Row 86 was an empty duplicate deviation row for @CGAP-LT-3 (case lives on row 84).
DELETE_ROWS = (82, 86, 87)


COVERAGE = "Automation_Coverage"
COVERAGE_FIRST_ROW = 2
STATES = ("yes", "partial", "deviation", "deferred", "no")


def refresh_coverage_row(book) -> None:
    """Retype this module's Automation_Coverage row (QA-owned sheet, typed values).

    Only the Lab Templates row and the TOTAL are touched; the other rows keep
    whatever QA has, so TOTAL stays consistent with what the sheet displays.
    """
    sheet = book[SHEET]
    counts = dict.fromkeys(STATES, 0)
    total = 0
    for row in range(2, sheet.max_row + 1):
        if not str(sheet.cell(row, 1).value or "").strip():
            continue
        total += 1
        state = str(sheet.cell(row, 11).value or "").strip().lower()
        if state in counts:
            counts[state] += 1

    coverage = book[COVERAGE]
    target = next(row for row in range(COVERAGE_FIRST_ROW, coverage.max_row + 1)
                  if str(coverage.cell(row, 1).value or "").strip() == SHEET)
    values = [counts[state] for state in STATES] + [total]
    for offset, value in enumerate(values, start=2):
        coverage.cell(target, offset).value = value
    coverage.cell(target, 8).value = f"{round(100 * counts['yes'] / total)}%" if total else "0%"

    total_row = next(row for row in range(COVERAGE_FIRST_ROW, coverage.max_row + 1)
                     if str(coverage.cell(row, 1).value or "").upper().startswith("TOTAL"))
    sums = [0] * 6
    for row in range(COVERAGE_FIRST_ROW, total_row):
        for index in range(6):
            value = coverage.cell(row, index + 2).value
            sums[index] += value if isinstance(value, (int, float)) else 0
    for index, value in enumerate(sums, start=2):
        coverage.cell(total_row, index).value = value
    coverage.cell(total_row, 8).value = f"{round(100 * sums[0] / sums[5])}%" if sums[5] else "0%"
    print(f"{COVERAGE} row {target}: {values} ({coverage.cell(target, 8).value}); "
          f"TOTAL now {sums}")


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sheet = book[SHEET]

    for row in sorted(DELETE_ROWS, reverse=True):
        title = str(sheet.cell(row, 1).value or "")
        summary = str(sheet.cell(row, 2).value or "").lower()
        if "bulk export" in summary or "re-imported" in summary or "Gap Check" in title:
            sheet.delete_rows(row, 1)
            print(f"deleted sheet row {row} (bulk export/import case)")
        elif row in (86, 87) and not title.strip() and not summary.strip():
            sheet.delete_rows(row, 1)
            print(f"deleted sheet row {row} (empty phantom case row)")
        else:
            print(f"skip delete row {row} — content changed ({title[:40]!r})")

    for row, reference in VERIFIED.items():
        sheet.cell(row, 11).value = "Yes"
        sheet.cell(row, 12).value = reference
        # The @labtemplates suite executed these cases green, which is what the
        # dashboard's Executed column counts (Automated=Yes AND a run status).
        sheet.cell(row, 7).value = "Pass"

    for row, reference in DEVIATIONS.items():
        sheet.cell(row, 11).value = "Deviation"
        sheet.cell(row, 12).value = reference
        sheet.cell(row, 7).value = "Fail"

    for row, reference in PARTIAL.items():
        sheet.cell(row, 11).value = "Partial"
        sheet.cell(row, 12).value = reference

    for row, reference in DEFERRED.items():
        sheet.cell(row, 11).value = "Deferred"
        sheet.cell(row, 12).value = reference
        if not str(sheet.cell(row, 7).value or "").strip():
            sheet.cell(row, 7).value = "NA"

    for row, reason in BLOCKED.items():
        sheet.cell(row, 11).value = "No"
        sheet.cell(row, 12).value = reason
        sheet.cell(row, 7).value = "NA"

    refresh_coverage_row(book)
    book.save(WORKBOOK)
    print(f"marked {len(VERIFIED)} Lab Templates cases Yes, {len(PARTIAL)} Partial, "
          f"{len(DEVIATIONS)} Deviations, {len(DEFERRED)} Deferred")


if __name__ == "__main__":
    main()
